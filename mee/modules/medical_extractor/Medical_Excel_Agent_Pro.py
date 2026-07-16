import sys
import json
import requests
import logging
from PyQt5.QtWidgets import (QApplication, QMainWindow, QWidget, QVBoxLayout,
                             QHBoxLayout, QPushButton, QTextEdit, QLabel,
                             QFileDialog, QTabWidget, QLineEdit, QComboBox,
                             QMessageBox, QTableWidget, QTableWidgetItem,
                             QSplitter, QGroupBox, QFormLayout, QSpinBox,
                             QDoubleSpinBox, QListWidget, QListWidgetItem,
                             QProgressBar, QHeaderView, QCheckBox, QDialog,
                             QDialogButtonBox, QScrollArea, QTextBrowser, QFrame,
                             QGraphicsDropShadowEffect)
from PyQt5.QtCore import Qt, QThread, pyqtSignal, QTimer, QPropertyAnimation, QEasingCurve, QRect, QSize, pyqtProperty
from PyQt5.QtGui import QFont, QColor, QPalette, QIcon, QPainter, QLinearGradient
import pandas as pd
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
import os
from datetime import datetime
import uuid
import traceback
from typing import Dict, List, Optional, Tuple
import time

# 抽取核心逻辑已抽离到 engine.py（供主流水线与本 GUI 共用）。
# 兼容两种运行方式：作为独立脚本（cwd 在本目录）或作为包模块导入。
try:
    from engine import (  # type: ignore
        APIClient, ClaudeClient, OpenAICompatibleClient,
        create_api_client, export_rows_to_excel,
    )
    from prompt_engineering_dialog import PromptEngineeringDialog
except ImportError:  # pragma: no cover - 包内导入路径
    from .engine import (
        APIClient, ClaudeClient, OpenAICompatibleClient,
        create_api_client, export_rows_to_excel,
    )
    from .prompt_engineering_dialog import PromptEngineeringDialog

# 兼容旧引用名
DeepSeekClient = OpenAICompatibleClient
OpenAIClient = OpenAICompatibleClient

# 配置日志
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(levelname)s - %(message)s',
    handlers=[
        logging.FileHandler('medical_extraction.log', encoding='utf-8'),
        logging.StreamHandler()
    ]
)
logger = logging.getLogger(__name__)


# ==================== 现代化UI组件 ====================

class ModernButton(QPushButton):
    """现代化动画按钮"""
    def __init__(self, text, parent=None, color='primary'):
        super().__init__(text, parent)
        self.setMinimumHeight(42)
        self.setCursor(Qt.PointingHandCursor)

        # 颜色方案
        self.colors = {
            'primary': ('#667eea', '#764ba2'),      # 紫色渐变
            'success': ('#56ab2f', '#a8e063'),      # 绿色渐变
            'danger': ('#eb3349', '#f45c43'),       # 红色渐变
            'warning': ('#f2994a', '#f2c94c'),      # 橙色渐变
            'info': ('#4facfe', '#00f2fe'),         # 蓝色渐变
            'dark': ('#2c3e50', '#34495e')          # 深色渐变
        }
        self.color_scheme = self.colors.get(color, self.colors['primary'])

        # 添加阴影效果
        shadow = QGraphicsDropShadowEffect(self)
        shadow.setBlurRadius(15)
        shadow.setColor(QColor(0, 0, 0, 80))
        shadow.setOffset(0, 3)
        self.setGraphicsEffect(shadow)

        # 设置样式
        self.update_style()

    def update_style(self):
        """更新按钮样式"""
        gradient_start = self.color_scheme[0]
        gradient_end = self.color_scheme[1]

        self.setStyleSheet(f"""
            QPushButton {{
                background: qlineargradient(x1:0, y1:0, x2:1, y2:0,
                    stop:0 {gradient_start}, stop:1 {gradient_end});
                color: white;
                border: none;
                border-radius: 21px;
                padding: 10px 24px;
                font-size: 14px;
                font-weight: 600;
                letter-spacing: 0.5px;
            }}
            QPushButton:hover {{
                background: qlineargradient(x1:0, y1:0, x2:1, y2:0,
                    stop:0 {gradient_end}, stop:1 {gradient_start});
                transform: translateY(-2px);
            }}
            QPushButton:pressed {{
                transform: translateY(0px);
            }}
            QPushButton:disabled {{
                background: #bdc3c7;
                color: #7f8c8d;
            }}
        """)

    def enterEvent(self, event):
        """鼠标进入动画"""
        self.animate_scale(1.05)
        super().enterEvent(event)

    def leaveEvent(self, event):
        """鼠标离开动画"""
        self.animate_scale(1.0)
        super().leaveEvent(event)

    def animate_scale(self, scale):
        """缩放动画"""
        self.animation = QPropertyAnimation(self, b"minimumHeight")
        self.animation.setDuration(200)
        self.animation.setEasingCurve(QEasingCurve.OutCubic)
        current_height = self.minimumHeight()
        target_height = int(42 * scale)
        self.animation.setStartValue(current_height)
        self.animation.setEndValue(target_height)
        self.animation.start()


class ModernCard(QFrame):
    """现代化卡片容器"""
    def __init__(self, parent=None):
        super().__init__(parent)
        self.setFrameShape(QFrame.StyledPanel)
        self.setStyleSheet("""
            QFrame {
                background-color: white;
                border-radius: 16px;
                border: 1px solid #e8e8e8;
                padding: 24px;
            }
        """)

        # 添加阴影
        shadow = QGraphicsDropShadowEffect(self)
        shadow.setBlurRadius(20)
        shadow.setColor(QColor(0, 0, 0, 30))
        shadow.setOffset(0, 4)
        self.setGraphicsEffect(shadow)


# ==================== 错误处理类 ====================

class ExtractionError:
    """提取错误记录"""
    def __init__(self, source: str, error_type: str, message: str, traceback: str = ""):
        self.source = source
        self.error_type = error_type
        self.message = message
        self.traceback = traceback
        self.timestamp = datetime.now()

    def to_dict(self):
        return {
            'source': self.source,
            'error_type': self.error_type,
            'message': self.message,
            'traceback': self.traceback,
            'timestamp': self.timestamp.isoformat()
        }


class ErrorManager:
    """错误管理器"""
    def __init__(self):
        self.errors: List[ExtractionError] = []

    def add_error(self, source: str, error_type: str, message: str, tb: str = ""):
        """添加错误记录"""
        error = ExtractionError(source, error_type, message, tb)
        self.errors.append(error)
        logger.error(f"[{source}] {error_type}: {message}")

    def get_error_summary(self) -> Dict:
        """获取错误摘要"""
        return {
            'total_errors': len(self.errors),
            'by_type': self._count_by_type(),
            'recent_errors': [e.to_dict() for e in self.errors[-10:]]
        }

    def _count_by_type(self) -> Dict[str, int]:
        """按类型统计错误"""
        counts = {}
        for error in self.errors:
            counts[error.error_type] = counts.get(error.error_type, 0) + 1
        return counts

    def export_errors(self, filepath: str):
        """导出错误报告"""
        try:
            with open(filepath, 'w', encoding='utf-8') as f:
                json.dump([e.to_dict() for e in self.errors], f, indent=2, ensure_ascii=False)
            logger.info(f"错误报告已导出到: {filepath}")
        except Exception as e:
            logger.error(f"导出错误报告失败: {str(e)}")

    def clear(self):
        """清空错误记录"""
        self.errors.clear()


# ==================== API客户端 ====================
# APIClient / OpenAICompatibleClient / ClaudeClient / create_api_client
# 均已迁移到 engine.py，本文件在顶部 import。



# ==================== 任务相关类 ====================

class Task:
    """任务类"""
    def __init__(self, name, template_path, output_path):
        self.id = str(uuid.uuid4())
        self.name = name
        self.template_path = template_path
        self.output_path = output_path
        self.emr_data = []  # 电子病历数据列表 [{source, content}]
        self.status = 'pending'
        self.progress = 0
        self.total = 0
        self.current_index = 0
        self.results = []
        self.created_time = datetime.now()
        self.error_msg = ''
        self.errors = []  # 错误列表
        self.success_count = 0
        self.failed_count = 0

    def add_emr_data(self, source: str, content: str):
        """添加电子病历数据"""
        self.emr_data.append({'source': source, 'content': content})
        self.total = len(self.emr_data)

    def to_dict(self):
        return {
            'id': self.id,
            'name': self.name,
            'template_path': self.template_path,
            'output_path': self.output_path,
            'emr_data': self.emr_data,
            'status': self.status,
            'progress': self.progress,
            'total': self.total,
            'current_index': self.current_index,
            'results': self.results,
            'created_time': self.created_time.isoformat(),
            'error_msg': self.error_msg,
            'errors': self.errors,
            'success_count': self.success_count,
            'failed_count': self.failed_count
        }

    @staticmethod
    def from_dict(data):
        task = Task(data['name'], data['template_path'], data['output_path'])
        task.id = data['id']
        task.emr_data = data['emr_data']
        task.status = data['status']
        task.progress = data['progress']
        task.total = data['total']
        task.current_index = data['current_index']
        task.results = data['results']
        task.created_time = datetime.fromisoformat(data['created_time'])
        task.error_msg = data.get('error_msg', '')
        task.errors = data.get('errors', [])
        task.success_count = data.get('success_count', 0)
        task.failed_count = data.get('failed_count', 0)
        return task


class TaskQueue:
    """任务队列管理器"""
    def __init__(self):
        self.tasks = []
        self.current_task = None

    def add_task(self, task):
        self.tasks.append(task)

    def remove_task(self, task_id):
        self.tasks = [t for t in self.tasks if t.id != task_id]
        if self.current_task and self.current_task.id == task_id:
            self.current_task = None

    def get_task(self, task_id):
        for task in self.tasks:
            if task.id == task_id:
                return task
        return None

    def get_next_task(self):
        for task in self.tasks:
            if task.status == 'pending':
                return task
        return None

    def save_to_file(self, filename='task_queue.json'):
        data = {
            'tasks': [task.to_dict() for task in self.tasks],
            'current_task_id': self.current_task.id if self.current_task else None
        }
        with open(filename, 'w', encoding='utf-8') as f:
            json.dump(data, f, indent=4, ensure_ascii=False)

    def load_from_file(self, filename='task_queue.json'):
        if not os.path.exists(filename):
            return

        with open(filename, 'r', encoding='utf-8') as f:
            data = json.load(f)

        self.tasks = [Task.from_dict(task_data) for task_data in data['tasks']]

        current_task_id = data.get('current_task_id')
        if current_task_id:
            self.current_task = self.get_task(current_task_id)


# ==================== 批量处理线程 ====================

class BatchProcessThread(QThread):
    """批量处理线程"""
    progress_updated = pyqtSignal(int, int, str)
    task_completed = pyqtSignal(str, list)
    task_failed = pyqtSignal(str, str)
    emr_processed = pyqtSignal(int, dict, bool)  # index, data, success
    error_occurred = pyqtSignal(str, str, str)  # source, error_type, message

    def __init__(self, task, api_config, template_config):
        super().__init__()
        self.task = task
        self.api_config = api_config
        self.template_config = template_config
        self.is_paused = False
        self.is_stopped = False
        self.error_manager = ErrorManager()

        # 创建API客户端
        try:
            self.api_client = create_api_client(api_config)
        except Exception as e:
            logger.error(f"创建API客户端失败: {str(e)}")
            self.api_client = None

    def run(self):
        """执行批量处理"""
        if not self.api_client:
            self.task.status = 'failed'
            self.task.error_msg = 'API客户端创建失败'
            self.task_failed.emit(self.task.id, 'API客户端创建失败')
            return

        try:
            self.task.status = 'processing'
            self.task.success_count = 0
            self.task.failed_count = 0

            for i in range(self.task.current_index, len(self.task.emr_data)):
                # 检查是否暂停或停止
                while self.is_paused and not self.is_stopped:
                    self.msleep(100)

                if self.is_stopped:
                    self.task.status = 'paused'
                    return

                emr_item = self.task.emr_data[i]
                source = emr_item['source']
                content = emr_item['content']

                self.task.current_index = i

                # 更新进度
                self.progress_updated.emit(
                    i + 1,
                    self.task.total,
                    f'正在处理: {source}'
                )

                # 处理单个病历
                success = False
                extracted_data = {}

                try:
                    extracted_data = self.process_single_emr(content)
                    extracted_data['_source'] = source
                    extracted_data['_status'] = 'success'

                    self.task.results.append(extracted_data)
                    self.task.success_count += 1
                    success = True

                    logger.info(f"成功处理: {source}")

                except Exception as e:
                    # 记录错误但继续处理下一个
                    error_type = type(e).__name__
                    error_msg = str(e)
                    tb = traceback.format_exc()

                    self.error_manager.add_error(source, error_type, error_msg, tb)
                    self.task.errors.append({
                        'source': source,
                        'error_type': error_type,
                        'message': error_msg,
                        'timestamp': datetime.now().isoformat()
                    })

                    # 创建错误记录
                    error_data = {
                        '_source': source,
                        '_status': 'failed',
                        '_error_type': error_type,
                        '_error_message': error_msg
                    }

                    self.task.results.append(error_data)
                    self.task.failed_count += 1

                    # 发送错误信号
                    self.error_occurred.emit(source, error_type, error_msg)

                    logger.error(f"处理失败 [{source}]: {error_msg}")

                # 更新进度
                self.task.progress = i + 1

                # 发送单个病历处理完成信号
                self.emr_processed.emit(i, extracted_data, success)

            # 所有病历处理完成
            self.task.status = 'completed'
            self.task.progress = self.task.total

            # 导出错误报告
            if self.task.errors:
                error_report_path = self.task.output_path.replace('.xlsx', '_errors.json')
                self.error_manager.export_errors(error_report_path)

            self.task_completed.emit(self.task.id, self.task.results)

        except Exception as e:
            # 致命错误
            tb = traceback.format_exc()
            logger.error(f"批量处理出现致命错误: {str(e)}\n{tb}")

            self.task.status = 'failed'
            self.task.error_msg = str(e)
            self.task_failed.emit(self.task.id, str(e))

    def process_single_emr(self, emr_content: str) -> Dict:
        """处理单个电子病历"""
        # 构建提示词
        prompt = self.build_prompt(emr_content)

        # 调用API
        api_response = self.api_client.call(prompt)

        # 解析响应
        extracted_data = self.parse_response(api_response)

        # 数据验证
        validated_data = self.validate_data(extracted_data)

        return validated_data

    def build_prompt(self, emr_content: str) -> str:
        """构建提示词"""
        fields_info = []
        for field in self.template_config['fields']:
            desc = field['description'] if field['description'] else field['column']
            fields_info.append(f"- {field['column']}: {desc} (类型: {field['type']})")

        fields_str = '\n'.join(fields_info)
        emr_format = self.template_config.get('emr_format', '')

        prompt = f"""你是一个专业的医疗数据提取助手。请从以下电子病历中提取信息，并按照指定的Excel模板字段填充数据。

电子病历格式说明：
{emr_format if emr_format else '标准电子病历格式'}

需要提取的字段：
{fields_str}

电子病历内容：
{emr_content}

请以JSON格式返回提取的数据，格式如下：
{{
    "字段1": "值1",
    "字段2": "值2",
    ...
}}

注意事项：
1. 严格按照字段的数据类型返回数据
2. 日期格式统一为 YYYY-MM-DD
3. 数字类型不要包含单位，只返回数字
4. 只返回JSON，不要包含其他解释文字
"""
        return prompt

    def parse_response(self, response: str) -> Dict:
        """解析API响应"""
        response = response.strip()

        # 移除可能的markdown代码块标记
        if response.startswith('```'):
            lines = response.split('\n')
            if len(lines) > 2:
                response = '\n'.join(lines[1:-1])
            if response.startswith('json'):
                response = response[4:].strip()

        try:
            data = json.loads(response)
            return data
        except json.JSONDecodeError as e:
            raise Exception(f'解析AI响应失败: {str(e)}\n原始响应: {response[:200]}...')

    def validate_data(self, data: Dict) -> Dict:
        """验证和清洗数据"""
        validated = {}

        for field in self.template_config['fields']:
            field_name = field['column']
            field_type = field.get('type', '文本')
            value = data.get(field_name, '-1')

            # 类型验证和转换
            try:
                if field_type in ['数字', '整数']:
                    if value == '-1' or value == -1:
                        validated[field_name] = -1
                    else:
                        validated[field_name] = int(float(str(value)))

                elif field_type in ['小数', '浮点数']:
                    if value == '-1' or value == -1:
                        validated[field_name] = -1
                    else:
                        validated[field_name] = float(str(value))

                else:
                    # 文本类型
                    if value is None or value == '':
                        validated[field_name] = '-1'
                    else:
                        validated[field_name] = str(value)

            except (ValueError, TypeError) as e:
                logger.warning(f"字段 {field_name} 类型转换失败: {value} -> {field_type}")
                validated[field_name] = '-1'

        return validated

    def pause(self):
        self.is_paused = True

    def resume(self):
        self.is_paused = False

    def stop(self):
        self.is_stopped = True


# ==================== UI组件 ====================

class ModernAPIConfigWidget(QWidget):
    """现代化API配置界面"""
    def __init__(self):
        super().__init__()
        self.init_ui()
        self.apply_styles()

    def init_ui(self):
        layout = QVBoxLayout()

        # 标题
        title = QLabel('🔧 API配置')
        title.setFont(QFont('Arial', 16, QFont.Bold))
        layout.addWidget(title)

        # 配置表单
        form_group = QGroupBox('基础配置')
        form_layout = QFormLayout()

        # API提供商选择
        self.provider_combo = QComboBox()
        self.provider_combo.addItems([
            'DeepSeek',
            'OpenAI',
            'Claude',
            'Azure OpenAI',
            '智谱AI',
            '通义千问',
            '自定义'
        ])
        self.provider_combo.currentTextChanged.connect(self.on_provider_changed)
        form_layout.addRow('AI提供商:', self.provider_combo)

        # API地址
        self.api_url_input = QLineEdit()
        self.api_url_input.setText('https://api.deepseek.com/v1/chat/completions')
        self.api_url_input.setPlaceholderText('请输入API地址')
        form_layout.addRow('API地址:', self.api_url_input)

        # API密钥
        self.api_key_input = QLineEdit()
        self.api_key_input.setEchoMode(QLineEdit.Password)
        self.api_key_input.setPlaceholderText('请输入API Key')

        # 显示/隐藏密钥按钮
        key_layout = QHBoxLayout()
        key_layout.addWidget(self.api_key_input)

        self.show_key_checkbox = QCheckBox('显示')
        self.show_key_checkbox.stateChanged.connect(self.toggle_key_visibility)
        key_layout.addWidget(self.show_key_checkbox)

        form_layout.addRow('API Key:', key_layout)

        # 模型名称
        self.model_input = QLineEdit()
        self.model_input.setText('deepseek-v4-pro')
        self.model_input.setPlaceholderText('请输入模型名称')
        form_layout.addRow('模型名称:', self.model_input)

        form_group.setLayout(form_layout)
        layout.addWidget(form_group)

        # 高级配置
        advanced_group = QGroupBox('高级配置')
        advanced_layout = QFormLayout()

        # 温度参数
        self.temperature_input = QDoubleSpinBox()
        self.temperature_input.setRange(0, 2)
        self.temperature_input.setValue(0.0)
        self.temperature_input.setSingleStep(0.1)
        self.temperature_input.setDecimals(1)
        advanced_layout.addRow('Temperature:', self.temperature_input)

        # 最大token数
        self.max_tokens_input = QSpinBox()
        self.max_tokens_input.setRange(100, 128000)
        self.max_tokens_input.setValue(65536)
        self.max_tokens_input.setSingleStep(1000)
        advanced_layout.addRow('最大Tokens:', self.max_tokens_input)

        # 超时时间
        self.timeout_input = QSpinBox()
        self.timeout_input.setRange(10, 600)
        self.timeout_input.setValue(180)
        self.timeout_input.setSuffix(' 秒')
        advanced_layout.addRow('请求超时:', self.timeout_input)

        # 重试次数
        self.retry_input = QSpinBox()
        self.retry_input.setRange(1, 10)
        self.retry_input.setValue(3)
        advanced_layout.addRow('重试次数:', self.retry_input)

        advanced_group.setLayout(advanced_layout)
        layout.addWidget(advanced_group)

        # 按钮组
        btn_layout = QHBoxLayout()

        test_btn = QPushButton('🔍 测试连接')
        test_btn.clicked.connect(self.test_connection)
        test_btn.setMinimumHeight(35)
        btn_layout.addWidget(test_btn)

        save_btn = QPushButton('💾 保存配置')
        save_btn.clicked.connect(self.save_config)
        save_btn.setMinimumHeight(35)
        btn_layout.addWidget(save_btn)

        load_btn = QPushButton('📂 加载配置')
        load_btn.clicked.connect(self.load_config)
        load_btn.setMinimumHeight(35)
        btn_layout.addWidget(load_btn)

        layout.addLayout(btn_layout)

        layout.addStretch()
        self.setLayout(layout)

        # 加载配置
        self.load_config()

    def apply_styles(self):
        """应用现代化样式"""
        self.setStyleSheet("""
            QGroupBox {
                font-weight: bold;
                border: 2px solid #3498db;
                border-radius: 5px;
                margin-top: 10px;
                padding-top: 10px;
            }
            QGroupBox::title {
                subcontrol-origin: margin;
                left: 10px;
                padding: 0 5px;
            }
            QPushButton {
                background-color: #3498db;
                color: white;
                border: none;
                border-radius: 5px;
                padding: 8px;
                font-weight: bold;
            }
            QPushButton:hover {
                background-color: #2980b9;
            }
            QPushButton:pressed {
                background-color: #21618c;
            }
            QLineEdit, QSpinBox, QDoubleSpinBox, QComboBox {
                padding: 5px;
                border: 1px solid #bdc3c7;
                border-radius: 3px;
            }
            QLineEdit:focus, QSpinBox:focus, QDoubleSpinBox:focus, QComboBox:focus {
                border: 2px solid #3498db;
            }
        """)

    def toggle_key_visibility(self, state):
        """切换密钥可见性"""
        if state == Qt.Checked:
            self.api_key_input.setEchoMode(QLineEdit.Normal)
        else:
            self.api_key_input.setEchoMode(QLineEdit.Password)

    def on_provider_changed(self, provider):
        """根据提供商更新默认配置"""
        configs = {
            'DeepSeek': {
                'url': 'https://api.deepseek.com/v1/chat/completions',
                'model': 'deepseek-v4-pro'
            },
            'OpenAI': {
                'url': 'https://api.openai.com/v1/chat/completions',
                'model': 'gpt-4'
            },
            'Claude': {
                'url': 'https://api.anthropic.com/v1/messages',
                'model': 'claude-3-opus-20240229'
            },
            'Azure OpenAI': {
                'url': 'https://YOUR_RESOURCE.openai.azure.com/openai/deployments/YOUR_DEPLOYMENT/chat/completions?api-version=2023-05-15',
                'model': 'gpt-4'
            },
            '智谱AI': {
                'url': 'https://open.bigmodel.cn/api/paas/v4/chat/completions',
                'model': 'glm-4'
            },
            '通义千问': {
                'url': 'https://dashscope.aliyuncs.com/api/v1/services/aigc/text-generation/generation',
                'model': 'qwen-max'
            }
        }

        if provider in configs:
            self.api_url_input.setText(configs[provider]['url'])
            self.model_input.setText(configs[provider]['model'])

    def get_config(self) -> Dict:
        """获取当前配置"""
        return {
            'provider': self.provider_combo.currentText(),
            'api_url': self.api_url_input.text(),
            'api_key': self.api_key_input.text(),
            'model': self.model_input.text(),
            'temperature': self.temperature_input.value(),
            'max_tokens': self.max_tokens_input.value(),
            'timeout': self.timeout_input.value(),
            'max_retries': self.retry_input.value()
        }

    def save_config(self):
        """保存配置"""
        config = self.get_config()
        try:
            with open('api_config.json', 'w', encoding='utf-8') as f:
                json.dump(config, f, indent=4, ensure_ascii=False)
            QMessageBox.information(self, '成功', '✅ API配置已保存')
            logger.info("API配置已保存")
        except Exception as e:
            QMessageBox.warning(self, '错误', f'❌ 保存配置失败: {str(e)}')
            logger.error(f"保存配置失败: {str(e)}")

    def load_config(self):
        """加载配置"""
        try:
            if os.path.exists('api_config.json'):
                with open('api_config.json', 'r', encoding='utf-8') as f:
                    config = json.load(f)

                self.provider_combo.setCurrentText(config.get('provider', 'DeepSeek'))
                self.api_url_input.setText(config.get('api_url', ''))
                self.api_key_input.setText(config.get('api_key', ''))
                self.model_input.setText(config.get('model', 'deepseek-v4-pro'))
                self.temperature_input.setValue(config.get('temperature', 0.0))
                self.max_tokens_input.setValue(config.get('max_tokens', 65536))
                self.timeout_input.setValue(config.get('timeout', 180))
                self.retry_input.setValue(config.get('max_retries', 3))

                logger.info("API配置已加载")
        except Exception as e:
            logger.warning(f"加载配置失败: {str(e)}")

    def test_connection(self):
        """测试API连接"""
        config = self.get_config()

        if not config['api_key']:
            QMessageBox.warning(self, '错误', '❌ 请先输入API Key')
            return

        try:
            # 创建API客户端
            client = create_api_client(config)

            # 测试连接
            success, message = client.test_connection()

            if success:
                QMessageBox.information(self, '成功', f'✅ {message}')
                logger.info(f"API连接测试成功: {config['provider']}")
            else:
                QMessageBox.warning(self, '失败', f'❌ {message}')
                logger.warning(f"API连接测试失败: {message}")

        except Exception as e:
            QMessageBox.critical(self, '错误', f'❌ 连接测试失败:\n{str(e)}')
            logger.error(f"API连接测试异常: {str(e)}")


# ==================== 模板配置界面 ====================

class TemplateConfigWidget(QWidget):
    """模板配置界面"""

    def __init__(self):
        super().__init__()
        self.template_path = None
        self.current_preset = None  # 当前预设名称
        self.init_ui()

    def init_ui(self):
        layout = QVBoxLayout()

        # 标题
        title = QLabel('📄 模板配置')
        title.setFont(QFont('Arial', 16, QFont.Bold))
        layout.addWidget(title)

        # 预设管理组
        preset_group = QGroupBox('预设管理')
        preset_layout = QVBoxLayout()

        # 预设列表
        preset_list_label = QLabel('我的预设:')
        preset_list_label.setFont(QFont('Arial', 12, QFont.Bold))
        preset_layout.addWidget(preset_list_label)

        self.preset_list = QListWidget()
        self.preset_list.setMinimumHeight(120)
        self.preset_list.setMaximumHeight(150)
        self.preset_list.itemSelectionChanged.connect(self.on_preset_selected)
        preset_layout.addWidget(self.preset_list)

        # 预设操作按钮 - 第一行
        preset_btn_layout1 = QHBoxLayout()

        new_preset_btn = QPushButton('➕ 新建预设')
        new_preset_btn.setMinimumHeight(38)
        new_preset_btn.clicked.connect(self.create_new_preset)
        preset_btn_layout1.addWidget(new_preset_btn)

        rename_preset_btn = QPushButton(' 重命名')
        rename_preset_btn.setMinimumHeight(38)
        rename_preset_btn.clicked.connect(self.rename_preset)
        preset_btn_layout1.addWidget(rename_preset_btn)

        delete_preset_btn = QPushButton('🗑️ 删除')
        delete_preset_btn.setMinimumHeight(38)
        delete_preset_btn.clicked.connect(self.delete_preset)
        preset_btn_layout1.addWidget(delete_preset_btn)

        preset_layout.addLayout(preset_btn_layout1)

        # 预设操作按钮 - 第二行
        preset_btn_layout2 = QHBoxLayout()

        save_preset_btn = QPushButton('💾 保存到当前预设')
        save_preset_btn.setMinimumHeight(38)
        save_preset_btn.clicked.connect(self.save_to_preset)
        preset_btn_layout2.addWidget(save_preset_btn)

        load_preset_btn = QPushButton('📂 加载当前预设')
        load_preset_btn.setMinimumHeight(38)
        load_preset_btn.clicked.connect(self.load_from_preset)
        preset_btn_layout2.addWidget(load_preset_btn)

        preset_layout.addLayout(preset_btn_layout2)

        # 预设状态显示
        self.preset_status_label = QLabel('当前预设: 未选择')
        self.preset_status_label.setStyleSheet('color: #64748b; padding: 8px; background-color: #f8fafc; border-radius: 6px;')
        preset_layout.addWidget(self.preset_status_label)

        preset_group.setLayout(preset_layout)
        layout.addWidget(preset_group)

        # Excel模板配置
        template_group = QGroupBox('Excel模板配置')
        template_layout = QVBoxLayout()

        template_file_layout = QHBoxLayout()
        self.template_path_label = QLabel('未选择模板文件')
        select_template_btn = QPushButton('选择Excel模板')
        select_template_btn.clicked.connect(self.select_template)
        template_file_layout.addWidget(self.template_path_label)
        template_file_layout.addWidget(select_template_btn)
        template_layout.addLayout(template_file_layout)

        # 模板字段映射
        self.template_table = QTableWidget()
        self.template_table.setColumnCount(3)
        self.template_table.setHorizontalHeaderLabels(['Excel列名', '字段描述', '数据类型'])
        template_layout.addWidget(QLabel('模板字段映射:'))
        template_layout.addWidget(self.template_table)

        refresh_btn = QPushButton('刷新模板字段')
        refresh_btn.clicked.connect(self.refresh_template_fields)
        template_layout.addWidget(refresh_btn)

        template_group.setLayout(template_layout)
        layout.addWidget(template_group)

        # 病历格式配置
        emr_group = QGroupBox('电子病历格式配置')
        emr_layout = QVBoxLayout()

        emr_layout.addWidget(QLabel('病历格式说明:'))
        self.emr_format_text = QTextEdit()
        self.emr_format_text.setPlaceholderText(
            '请描述您的电子病历格式，例如：\n'
            '- 患者信息部分包含：姓名、性别、年龄、病历号\n'
            '- 主诉部分格式：主诉:xxxxx\n'
            '- 诊断部分格式：诊断:xxxxx\n'
            '- 用药部分格式：处方:xxxxx'
        )
        self.emr_format_text.setMaximumHeight(150)
        emr_layout.addWidget(self.emr_format_text)

        # 保存/加载格式
        btn_layout = QHBoxLayout()
        save_format_btn = QPushButton('保存格式配置')
        save_format_btn.clicked.connect(self.save_format_config)
        load_format_btn = QPushButton('加载格式配置')
        load_format_btn.clicked.connect(self.load_format_config)
        btn_layout.addWidget(save_format_btn)
        btn_layout.addWidget(load_format_btn)
        emr_layout.addLayout(btn_layout)

        emr_group.setLayout(emr_layout)
        layout.addWidget(emr_group)

        # AI提示词工程
        prompt_group = QGroupBox('AI提示词工程')
        prompt_layout = QVBoxLayout()

        prompt_info = QLabel('为每个字段生成智能提取规则，提升提取准确率')
        prompt_info.setWordWrap(True)
        prompt_info.setStyleSheet('color: #64748b; padding: 12px; background-color: #f8fafc; border-radius: 8px;')
        prompt_layout.addWidget(prompt_info)

        prompt_btn = QPushButton('打开提示词工程')
        prompt_btn.clicked.connect(self.open_prompt_engineering)
        prompt_btn.setMinimumHeight(40)
        prompt_layout.addWidget(prompt_btn)

        prompt_group.setLayout(prompt_layout)
        layout.addWidget(prompt_group)

        self.setLayout(layout)

        # 初始化预设列表
        self.load_preset_list()
        self.load_format_config()

    def select_template(self):
        """选择Excel模板文件"""
        file_path, _ = QFileDialog.getOpenFileName(
            self, '选择Excel模板', '', 'Excel Files (*.xlsx *.xls)'
        )
        if file_path:
            self.template_path = file_path
            self.template_path_label.setText(os.path.basename(file_path))
            self.refresh_template_fields()

    def refresh_template_fields(self):
        """刷新模板字段"""
        if not self.template_path:
            QMessageBox.warning(self, '提示', '请先选择Excel模板')
            return

        try:
            df = pd.read_excel(self.template_path, nrows=0)
            columns = df.columns.tolist()

            self.template_table.setRowCount(len(columns))
            for i, col in enumerate(columns):
                self.template_table.setItem(i, 0, QTableWidgetItem(col))
                self.template_table.setItem(i, 1, QTableWidgetItem(''))
                self.template_table.setItem(i, 2, QTableWidgetItem('文本'))

            QMessageBox.information(self, '成功', f'✅ 已读取 {len(columns)} 个字段')
            logger.info(f"加载模板字段: {len(columns)}个")
        except Exception as e:
            QMessageBox.warning(self, '错误', f'❌ 读取模板失败: {str(e)}')
            logger.error(f"读取模板失败: {str(e)}")

    def get_template_config(self):
        """获取模板配置"""
        fields = []
        for i in range(self.template_table.rowCount()):
            col_name = self.template_table.item(i, 0)
            description = self.template_table.item(i, 1)
            data_type = self.template_table.item(i, 2)

            if col_name:
                fields.append({
                    'column': col_name.text(),
                    'description': description.text() if description else '',
                    'type': data_type.text() if data_type else '文本'
                })

        return {
            'template_path': self.template_path,
            'fields': fields,
            'emr_format': self.emr_format_text.toPlainText()
        }

    def save_format_config(self):
        """保存格式配置"""
        config = self.get_template_config()
        try:
            with open('template_config.json', 'w', encoding='utf-8') as f:
                json.dump(config, f, indent=4, ensure_ascii=False)
            QMessageBox.information(self, '成功', '✅ 格式配置已保存')
            logger.info("模板配置已保存")
        except Exception as e:
            QMessageBox.warning(self, '错误', f'❌ 保存配置失败: {str(e)}')
            logger.error(f"保存配置失败: {str(e)}")

    def load_format_config(self):
        """加载格式配置"""
        try:
            if os.path.exists('template_config.json'):
                with open('template_config.json', 'r', encoding='utf-8') as f:
                    config = json.load(f)

                if config.get('template_path') and os.path.exists(config['template_path']):
                    self.template_path = config['template_path']
                    self.template_path_label.setText(os.path.basename(self.template_path))

                fields = config.get('fields', [])
                self.template_table.setRowCount(len(fields))
                for i, field in enumerate(fields):
                    self.template_table.setItem(i, 0, QTableWidgetItem(field.get('column', '')))
                    self.template_table.setItem(i, 1, QTableWidgetItem(field.get('description', '')))
                    self.template_table.setItem(i, 2, QTableWidgetItem(field.get('type', '文本')))

                self.emr_format_text.setPlainText(config.get('emr_format', ''))
                logger.info("模板配置已加载")
        except Exception as e:
            logger.warning(f"加载配置失败: {str(e)}")

    def open_prompt_engineering(self):
        """打开AI提示词工程对话框"""
        # 获取当前模板配置
        template_config = self.get_template_config()

        if not template_config['fields']:
            QMessageBox.warning(self, '提示', '请先选择Excel模板并刷新字段')
            return

        # 打开提示词工程对话框
        dialog = PromptEngineeringDialog(template_config, self)
        if dialog.exec_() == QDialog.Accepted:
            prompts = dialog.get_prompts()

            # 保存提示词配置
            try:
                with open('prompt_config.json', 'w', encoding='utf-8') as f:
                    json.dump(prompts, f, indent=4, ensure_ascii=False)

                QMessageBox.information(self, '成功', f'✅ 已为 {len(prompts)} 个字段保存提示词配置')
                logger.info(f"提示词配置已保存: {len(prompts)}个字段")
            except Exception as e:
                QMessageBox.warning(self, '错误', f'保存提示词配置失败:\n{str(e)}')
                logger.error(f"保存提示词配置失败: {str(e)}")

    # ==================== 预设管理方法 ====================

    def get_preset_filename(self, preset_name: str) -> str:
        """获取预设文件名"""
        # 使用安全的文件名（移除特殊字符）
        safe_name = "".join(c for c in preset_name if c.isalnum() or c in (' ', '-', '_')).strip()
        return f'template_preset_{safe_name}.json'

    def get_all_presets(self) -> list:
        """获取所有预设文件"""
        import glob
        preset_files = glob.glob('template_preset_*.json')
        presets = []
        for file in preset_files:
            # 从文件名中提取预设名称
            name = file.replace('template_preset_', '').replace('.json', '')
            presets.append(name)
        return sorted(presets)

    def load_preset_list(self):
        """加载预设列表"""
        self.preset_list.clear()
        presets = self.get_all_presets()

        if not presets:
            # 如果没有预设，创建一个默认预设
            self.preset_list.addItem('默认预设')
        else:
            for preset in presets:
                self.preset_list.addItem(preset)

        # 默认选择第一个
        if self.preset_list.count() > 0:
            self.preset_list.setCurrentRow(0)

    def on_preset_selected(self):
        """预设选择改变"""
        current_item = self.preset_list.currentItem()
        if current_item:
            self.current_preset = current_item.text()
            self.update_preset_status()
        else:
            self.current_preset = None
            self.preset_status_label.setText('当前预设: 未选择')

    def create_new_preset(self):
        """创建新预设"""
        from PyQt5.QtWidgets import QInputDialog

        preset_name, ok = QInputDialog.getText(
            self, '新建预设',
            '请输入预设名称:',
            QLineEdit.Normal,
            f'预设{self.preset_list.count() + 1}'
        )

        if ok and preset_name:
            preset_name = preset_name.strip()

            if not preset_name:
                QMessageBox.warning(self, '错误', '预设名称不能为空')
                return

            # 检查是否已存在
            existing_presets = self.get_all_presets()
            if preset_name in existing_presets:
                QMessageBox.warning(self, '错误', f'预设 "{preset_name}" 已存在')
                return

            # 添加到列表
            self.preset_list.addItem(preset_name)
            self.preset_list.setCurrentRow(self.preset_list.count() - 1)

            QMessageBox.information(self, '成功', f'✅ 已创建预设 "{preset_name}"\n请配置模板后保存')
            logger.info(f"创建新预设: {preset_name}")

    def rename_preset(self):
        """重命名预设"""
        from PyQt5.QtWidgets import QInputDialog

        current_item = self.preset_list.currentItem()
        if not current_item:
            QMessageBox.warning(self, '提示', '请先选择要重命名的预设')
            return

        old_name = current_item.text()
        old_file = self.get_preset_filename(old_name)

        new_name, ok = QInputDialog.getText(
            self, '重命名预设',
            '请输入新的预设名称:',
            QLineEdit.Normal,
            old_name
        )

        if ok and new_name:
            new_name = new_name.strip()

            if not new_name:
                QMessageBox.warning(self, '错误', '预设名称不能为空')
                return

            if new_name == old_name:
                return

            # 检查新名称是否已存在
            existing_presets = self.get_all_presets()
            if new_name in existing_presets:
                QMessageBox.warning(self, '错误', f'预设 "{new_name}" 已存在')
                return

            new_file = self.get_preset_filename(new_name)

            # 如果旧预设文件存在，则重命名文件
            if os.path.exists(old_file):
                try:
                    os.rename(old_file, new_file)
                    logger.info(f"重命名预设文件: {old_file} -> {new_file}")
                except Exception as e:
                    QMessageBox.warning(self, '错误', f'重命名文件失败:\n{str(e)}')
                    return

            # 更新列表
            current_item.setText(new_name)
            self.current_preset = new_name
            self.update_preset_status()

            QMessageBox.information(self, '成功', f'✅ 预设已重命名为 "{new_name}"')
            logger.info(f"重命名预设: {old_name} -> {new_name}")

    def delete_preset(self):
        """删除预设"""
        current_item = self.preset_list.currentItem()
        if not current_item:
            QMessageBox.warning(self, '提示', '请先选择要删除的预设')
            return

        preset_name = current_item.text()
        preset_file = self.get_preset_filename(preset_name)

        reply = QMessageBox.question(
            self, '确认删除',
            f'确定要删除预设 "{preset_name}" 吗？\n此操作不可恢复！',
            QMessageBox.Yes | QMessageBox.No
        )

        if reply == QMessageBox.Yes:
            # 删除文件（如果存在）
            if os.path.exists(preset_file):
                try:
                    os.remove(preset_file)
                    logger.info(f"删除预设文件: {preset_file}")
                except Exception as e:
                    QMessageBox.warning(self, '错误', f'删除文件失败:\n{str(e)}')
                    return

            # 从列表中移除
            row = self.preset_list.currentRow()
            self.preset_list.takeItem(row)

            # 如果列表为空，添加默认预设
            if self.preset_list.count() == 0:
                self.preset_list.addItem('默认预设')
                self.preset_list.setCurrentRow(0)

            QMessageBox.information(self, '成功', f'✅ 已删除预设 "{preset_name}"')
            logger.info(f"删除预设: {preset_name}")

    def on_preset_changed(self, preset_name):
        """预设选择改变（保留兼容性）"""
        self.current_preset = preset_name
        self.update_preset_status()

    def save_to_preset(self):
        """保存到当前预设"""
        if not self.current_preset:
            QMessageBox.warning(self, '提示', '请先选择一个预设')
            return

        config = self.get_template_config()

        if not config['fields']:
            QMessageBox.warning(self, '提示', '当前没有配置可保存，请先选择模板并配置字段')
            return

        preset_file = self.get_preset_filename(self.current_preset)

        try:
            with open(preset_file, 'w', encoding='utf-8') as f:
                json.dump(config, f, indent=4, ensure_ascii=False)

            self.update_preset_status()
            QMessageBox.information(self, '成功', f'✅ 配置已保存到 {self.current_preset}')
            logger.info(f"配置已保存到预设: {self.current_preset}")
        except Exception as e:
            QMessageBox.warning(self, '错误', f'❌ 保存预设失败: {str(e)}')
            logger.error(f"保存预设失败: {str(e)}")

    def load_from_preset(self):
        """加载当前预设"""
        if not self.current_preset:
            QMessageBox.warning(self, '提示', '请先选择一个预设')
            return

        preset_file = self.get_preset_filename(self.current_preset)

        if not os.path.exists(preset_file):
            QMessageBox.information(self, '提示', f'预设 {self.current_preset} 还没有保存过配置')
            return

        try:
            with open(preset_file, 'r', encoding='utf-8') as f:
                config = json.load(f)

            # 加载模板路径
            if config.get('template_path') and os.path.exists(config['template_path']):
                self.template_path = config['template_path']
                self.template_path_label.setText(os.path.basename(self.template_path))
            else:
                self.template_path = None
                self.template_path_label.setText('未选择模板')

            # 加载字段配置
            fields = config.get('fields', [])
            self.template_table.setRowCount(len(fields))
            for i, field in enumerate(fields):
                self.template_table.setItem(i, 0, QTableWidgetItem(field.get('column', '')))
                self.template_table.setItem(i, 1, QTableWidgetItem(field.get('description', '')))
                self.template_table.setItem(i, 2, QTableWidgetItem(field.get('type', '文本')))

            # 加载病历格式配置
            self.emr_format_text.setPlainText(config.get('emr_format', ''))

            self.update_preset_status()
            QMessageBox.information(self, '成功', f'✅ 已从 {self.current_preset} 加载配置\n共 {len(fields)} 个字段')
            logger.info(f"从预设加载配置: {self.current_preset}, {len(fields)}个字段")
        except Exception as e:
            QMessageBox.warning(self, '错误', f'❌ 加载预设失败: {str(e)}')
            logger.error(f"加载预设失败: {str(e)}")

    def update_preset_status(self):
        """更新预设状态显示"""
        if not self.current_preset:
            self.preset_status_label.setText('当前预设: 未选择')
            return

        preset_file = self.get_preset_filename(self.current_preset)

        if os.path.exists(preset_file):
            try:
                with open(preset_file, 'r', encoding='utf-8') as f:
                    config = json.load(f)
                field_count = len(config.get('fields', []))
                self.preset_status_label.setText(
                    f'当前预设: {self.current_preset} ✅ (已保存，{field_count}个字段)'
                )
                self.preset_status_label.setStyleSheet(
                    'color: #16a34a; padding: 8px; background-color: #dcfce7; border-radius: 6px; font-weight: bold;'
                )
            except:
                self.preset_status_label.setText(f'当前预设: {self.current_preset} ⚠️ (文件损坏)')
                self.preset_status_label.setStyleSheet(
                    'color: #dc2626; padding: 8px; background-color: #fee2e2; border-radius: 6px;'
                )
        else:
            self.preset_status_label.setText(f'当前预设: {self.current_preset} ⭕ (未保存)')
            self.preset_status_label.setStyleSheet(
                'color: #64748b; padding: 8px; background-color: #f8fafc; border-radius: 6px;'
            )


# ==================== Excel导入对话框 ====================

class ExcelImportDialog(QDialog):
    """改进的Excel导入对话框 - 支持病例勾选"""
    def __init__(self, parent=None):
        super().__init__(parent)
        self.setWindowTitle('从Excel导入病历')
        self.setMinimumSize(1000, 700)
        self.excel_file = None
        self.df = None
        self.selected_columns = []
        self.selected_rows = []  # 新增：选中的行索引
        self.init_ui()
        self.apply_styles()

    def init_ui(self):
        layout = QVBoxLayout()

        # 标题
        title = QLabel('从Excel表格批量导入病历')
        title.setFont(QFont('SF Pro Display', 18, QFont.Bold))
        title.setStyleSheet('color: #1e293b; padding: 10px 0;')
        layout.addWidget(title)

        # 文件选择
        file_group = QGroupBox('1. 选择Excel文件')
        file_layout = QHBoxLayout()

        self.file_label = QLabel('未选择文件')
        file_layout.addWidget(self.file_label)

        select_btn = QPushButton('📂 选择Excel文件')
        select_btn.clicked.connect(self.select_excel_file)
        file_layout.addWidget(select_btn)

        file_group.setLayout(file_layout)
        layout.addWidget(file_group)

        # 工作表选择
        sheet_group = QGroupBox('2. 选择工作表')
        sheet_layout = QHBoxLayout()

        self.sheet_combo = QComboBox()
        self.sheet_combo.currentTextChanged.connect(self.on_sheet_changed)
        sheet_layout.addWidget(QLabel('工作表:'))
        sheet_layout.addWidget(self.sheet_combo)

        sheet_group.setLayout(sheet_layout)
        layout.addWidget(sheet_group)

        # 列选择
        column_group = QGroupBox('3. 选择包含病历内容的列')
        column_layout = QVBoxLayout()

        tip_label = QLabel('提示: 每一行代表一个病人的病历，可选择多列组合')
        tip_label.setStyleSheet('color: #64748b; padding: 8px; background-color: #f8fafc; border-radius: 6px;')
        column_layout.addWidget(tip_label)

        self.column_list = QListWidget()
        self.column_list.setSelectionMode(QListWidget.MultiSelection)
        column_layout.addWidget(self.column_list)

        column_group.setLayout(column_layout)
        layout.addWidget(column_group)

        # 预览和选择
        preview_group = QGroupBox('4. 预览并选择要导入的病例')
        preview_layout = QVBoxLayout()

        # 操作按钮
        preview_btn_layout = QHBoxLayout()

        preview_btn = QPushButton('预览数据')
        preview_btn.setMinimumHeight(36)
        preview_btn.setCursor(Qt.PointingHandCursor)
        preview_btn.clicked.connect(self.preview_data)
        preview_btn_layout.addWidget(preview_btn)

        select_all_btn = QPushButton('全选')
        select_all_btn.setMinimumHeight(36)
        select_all_btn.setCursor(Qt.PointingHandCursor)
        select_all_btn.clicked.connect(self.select_all_rows)
        preview_btn_layout.addWidget(select_all_btn)

        deselect_all_btn = QPushButton('全不选')
        deselect_all_btn.setMinimumHeight(36)
        deselect_all_btn.setCursor(Qt.PointingHandCursor)
        deselect_all_btn.clicked.connect(self.deselect_all_rows)
        preview_btn_layout.addWidget(deselect_all_btn)

        invert_selection_btn = QPushButton('反选')
        invert_selection_btn.setMinimumHeight(36)
        invert_selection_btn.setCursor(Qt.PointingHandCursor)
        invert_selection_btn.clicked.connect(self.invert_selection)
        preview_btn_layout.addWidget(invert_selection_btn)

        preview_btn_layout.addStretch()
        preview_layout.addLayout(preview_btn_layout)

        # 预览表格（带复选框）
        self.preview_table = QTableWidget()
        self.preview_table.setSelectionMode(QTableWidget.MultiSelection)
        self.preview_table.setSelectionBehavior(QTableWidget.SelectRows)
        preview_layout.addWidget(self.preview_table)

        preview_group.setLayout(preview_layout)
        layout.addWidget(preview_group)

        # 分批设置
        batch_group = QGroupBox('5. 分批设置（可选）')
        batch_layout = QHBoxLayout()

        self.enable_batch_checkbox = QCheckBox('启用分批创建任务')
        self.enable_batch_checkbox.stateChanged.connect(self.on_batch_enabled_changed)
        batch_layout.addWidget(self.enable_batch_checkbox)

        batch_layout.addWidget(QLabel('每'))
        self.batch_size_spin = QSpinBox()
        self.batch_size_spin.setRange(1, 1000)
        self.batch_size_spin.setValue(50)
        self.batch_size_spin.setEnabled(False)
        batch_layout.addWidget(self.batch_size_spin)
        batch_layout.addWidget(QLabel('条病例创建一个任务'))

        batch_layout.addStretch()
        batch_group.setLayout(batch_layout)
        layout.addWidget(batch_group)

        # 统计信息
        self.stats_label = QLabel('待导入病历数: 0 | 已选择: 0')
        self.stats_label.setFont(QFont('Arial', 10, QFont.Bold))
        layout.addWidget(self.stats_label)

        # 按钮
        button_box = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel)
        button_box.button(QDialogButtonBox.Ok).setText('✅ 确认导入')
        button_box.button(QDialogButtonBox.Cancel).setText('❌ 取消')
        button_box.accepted.connect(self.validate_and_accept)
        button_box.rejected.connect(self.reject)
        layout.addWidget(button_box)

        self.setLayout(layout)

    def apply_styles(self):
        self.setStyleSheet("""
            QGroupBox {
                font-weight: bold;
                border: 2px solid #2ecc71;
                border-radius: 5px;
                margin-top: 10px;
                padding-top: 10px;
            }
            QPushButton {
                background-color: #27ae60;
                color: white;
                border: none;
                border-radius: 5px;
                padding: 8px;
                font-weight: bold;
                min-height: 30px;
                min-width: 80px;
            }
            QPushButton:hover {
                background-color: #229954;
            }
        """)

    def on_batch_enabled_changed(self, state):
        """分批选项改变"""
        self.batch_size_spin.setEnabled(state == Qt.Checked)

    def select_excel_file(self):
        """选择Excel文件"""
        file_path, _ = QFileDialog.getOpenFileName(
            self,
            '选择Excel文件',
            '',
            'Excel Files (*.xlsx *.xls)'
        )

        if file_path:
            try:
                self.excel_file = file_path
                self.file_label.setText(os.path.basename(file_path))

                # 读取工作表名称
                xl_file = pd.ExcelFile(file_path)
                self.sheet_combo.clear()
                self.sheet_combo.addItems(xl_file.sheet_names)

                logger.info(f"成功加载Excel文件: {file_path}")

            except Exception as e:
                QMessageBox.warning(self, '错误', f'读取Excel文件失败:\n{str(e)}')
                logger.error(f"读取Excel失败: {str(e)}")

    def on_sheet_changed(self, sheet_name):
        """工作表改变时"""
        if not self.excel_file or not sheet_name:
            return

        try:
            # 读取完整数据（不要用nrows限制）
            self.df = pd.read_excel(self.excel_file, sheet_name=sheet_name)

            # 更新列列表
            self.column_list.clear()
            for col in self.df.columns:
                self.column_list.addItem(str(col))

            logger.info(f"加载工作表: {sheet_name}, 行数: {len(self.df)}, 列数: {len(self.df.columns)}")

        except Exception as e:
            QMessageBox.warning(self, '错误', f'读取工作表失败:\n{str(e)}')
            logger.error(f"读取工作表失败: {str(e)}")

    def preview_data(self):
        """预览数据并允许选择"""
        if self.df is None:
            QMessageBox.warning(self, '提示', '请先选择Excel文件和工作表')
            return

        # 获取选中的列
        selected_items = self.column_list.selectedItems()
        if not selected_items:
            QMessageBox.warning(self, '提示', '请至少选择一列')
            return

        self.selected_columns = [item.text() for item in selected_items]

        # 显示所有数据（不只是前5行）
        display_df = self.df[self.selected_columns]

        # 设置表格
        self.preview_table.setRowCount(len(display_df))
        self.preview_table.setColumnCount(len(self.selected_columns) + 1)  # +1 for checkbox column

        headers = ['选择'] + self.selected_columns
        self.preview_table.setHorizontalHeaderLabels(headers)

        # 填充数据
        for i in range(len(display_df)):
            # 复选框列
            checkbox = QCheckBox()
            checkbox.setChecked(True)  # 默认全选
            checkbox_widget = QWidget()
            checkbox_layout = QHBoxLayout(checkbox_widget)
            checkbox_layout.addWidget(checkbox)
            checkbox_layout.setAlignment(Qt.AlignCenter)
            checkbox_layout.setContentsMargins(0, 0, 0, 0)
            self.preview_table.setCellWidget(i, 0, checkbox_widget)

            # 数据列
            for j, col in enumerate(self.selected_columns):
                value = str(display_df.iloc[i][col])
                self.preview_table.setItem(i, j + 1, QTableWidgetItem(value))

        self.preview_table.resizeColumnsToContents()

        # 更新统计
        total_rows = len(self.df)
        selected_count = self.count_selected_rows()
        self.stats_label.setText(f'✅ 待导入病历数: {total_rows} | 已选择: {selected_count}')

    def count_selected_rows(self) -> int:
        """统计选中的行数"""
        count = 0
        for i in range(self.preview_table.rowCount()):
            checkbox_widget = self.preview_table.cellWidget(i, 0)
            if checkbox_widget:
                checkbox = checkbox_widget.findChild(QCheckBox)
                if checkbox and checkbox.isChecked():
                    count += 1
        return count

    def select_all_rows(self):
        """全选"""
        for i in range(self.preview_table.rowCount()):
            checkbox_widget = self.preview_table.cellWidget(i, 0)
            if checkbox_widget:
                checkbox = checkbox_widget.findChild(QCheckBox)
                if checkbox:
                    checkbox.setChecked(True)
        self.update_stats()

    def deselect_all_rows(self):
        """全不选"""
        for i in range(self.preview_table.rowCount()):
            checkbox_widget = self.preview_table.cellWidget(i, 0)
            if checkbox_widget:
                checkbox = checkbox_widget.findChild(QCheckBox)
                if checkbox:
                    checkbox.setChecked(False)
        self.update_stats()

    def invert_selection(self):
        """反选"""
        for i in range(self.preview_table.rowCount()):
            checkbox_widget = self.preview_table.cellWidget(i, 0)
            if checkbox_widget:
                checkbox = checkbox_widget.findChild(QCheckBox)
                if checkbox:
                    checkbox.setChecked(not checkbox.isChecked())
        self.update_stats()

    def update_stats(self):
        """更新统计信息"""
        if self.df is not None:
            total_rows = len(self.df)
            selected_count = self.count_selected_rows()
            self.stats_label.setText(f'✅ 待导入病历数: {total_rows} | 已选择: {selected_count}')

    def validate_and_accept(self):
        """验证并接受"""
        if self.df is None or not self.selected_columns:
            QMessageBox.warning(self, '错误', '请先预览数据')
            return

        selected_count = self.count_selected_rows()
        if selected_count == 0:
            QMessageBox.warning(self, '错误', '请至少选择一条病例')
            return

        self.accept()

    def get_emr_data(self) -> List[Dict]:
        """获取选中的病历数据"""
        if self.df is None or not self.selected_columns:
            return []

        emr_data = []
        empty_rows = []  # 记录空行

        # 只获取选中的行
        for i in range(self.preview_table.rowCount()):
            checkbox_widget = self.preview_table.cellWidget(i, 0)
            if checkbox_widget:
                checkbox = checkbox_widget.findChild(QCheckBox)
                if checkbox and checkbox.isChecked():
                    # 获取该行数据
                    row = self.df.iloc[i]

                    # 组合选中的列
                    content_parts = []
                    for col in self.selected_columns:
                        value = row[col]
                        if pd.notna(value):
                            content_parts.append(f"{col}: {value}")

                    content = '\n'.join(content_parts)

                    if content.strip():
                        emr_data.append({
                            'source': f"Excel第{i + 2}行",  # +2因为Excel从1开始且有表头
                            'content': content,
                            'row_index': i  # 保存行索引用于错误追踪
                        })
                    else:
                        # 记录空行
                        empty_rows.append(i + 2)  # Excel行号

        # 如果有空行被跳过，给出警告
        if empty_rows:
            logger.warning(f"跳过了{len(empty_rows)}行空数据: Excel第{empty_rows}行")
            QMessageBox.warning(
                self, '提示',
                f'⚠️ 发现{len(empty_rows)}行数据在所选列中全部为空，已自动跳过：\n'
                f'Excel行号: {", ".join(map(str, empty_rows[:10]))}'
                f'{"..." if len(empty_rows) > 10 else ""}\n\n'
                f'实际导入: {len(emr_data)}行\n'
                f'已跳过: {len(empty_rows)}行'
            )

        return emr_data

    def is_batch_enabled(self) -> bool:
        """是否启用分批"""
        return self.enable_batch_checkbox.isChecked()

    def get_batch_size(self) -> int:
        """获取批次大小"""
        return self.batch_size_spin.value()

# ==================== 任务创建对话框（改进版） ====================

class ModernNewTaskDialog(QDialog):
    """现代化任务创建对话框"""
    def __init__(self, parent):
        super().__init__(parent)
        self.parent_window = parent
        self.template_path = None
        self.output_path = None
        self.emr_data = []
        self.excel_dialog = None  # 保存Excel导入对话框引用
        self.setup_ui()
        self.apply_styles()

    def setup_ui(self):
        self.setWindowTitle('📋 创建新任务')
        self.setMinimumWidth(700)
        self.setMinimumHeight(600)

        layout = QVBoxLayout()

        # 标题
        title = QLabel('创建数据提取任务')
        title.setFont(QFont('Arial', 16, QFont.Bold))
        layout.addWidget(title)

        # 任务信息
        info_group = QGroupBox('📝 任务信息')
        info_layout = QFormLayout()

        self.name_input = QLineEdit()
        self.name_input.setText(f'任务_{datetime.now().strftime("%Y%m%d_%H%M%S")}')
        self.name_input.setPlaceholderText('请输入任务名称')
        info_layout.addRow('任务名称:', self.name_input)

        info_group.setLayout(info_layout)
        layout.addWidget(info_group)

        # 模板配置
        template_group = QGroupBox('📄 模板配置')
        template_layout = QVBoxLayout()

        template_row = QHBoxLayout()
        self.template_label = QLabel('未选择模板')
        template_btn = QPushButton('选择Excel模板')
        template_btn.clicked.connect(self.select_template)
        template_row.addWidget(self.template_label)
        template_row.addWidget(template_btn)
        template_layout.addLayout(template_row)

        output_row = QHBoxLayout()
        self.output_label = QLabel('未选择输出路径')
        output_btn = QPushButton('选择输出路径')
        output_btn.clicked.connect(self.select_output)
        output_row.addWidget(self.output_label)
        output_row.addWidget(output_btn)
        template_layout.addLayout(output_row)

        template_group.setLayout(template_layout)
        layout.addWidget(template_group)

        # 病历导入
        emr_group = QGroupBox('🏥 病历导入')
        emr_layout = QVBoxLayout()

        # 导入方式选择
        import_tabs = QTabWidget()

        # Tab 1: 从文件导入
        file_tab = QWidget()
        file_layout = QVBoxLayout()

        self.file_list = QListWidget()
        self.file_list.setMinimumHeight(150)
        file_layout.addWidget(self.file_list)

        file_btn_layout = QHBoxLayout()
        add_files_btn = QPushButton('➕ 添加文件')
        add_files_btn.clicked.connect(self.add_emr_files)
        add_folder_btn = QPushButton('📁 添加文件夹')
        add_folder_btn.clicked.connect(self.add_emr_folder)
        clear_files_btn = QPushButton('🗑️ 清空')
        clear_files_btn.clicked.connect(self.clear_file_list)
        file_btn_layout.addWidget(add_files_btn)
        file_btn_layout.addWidget(add_folder_btn)
        file_btn_layout.addWidget(clear_files_btn)
        file_layout.addLayout(file_btn_layout)

        file_tab.setLayout(file_layout)
        import_tabs.addTab(file_tab, '📄 从文件导入')

        # Tab 2: 从Excel导入
        excel_tab = QWidget()
        excel_layout = QVBoxLayout()

        excel_info = QLabel('💡 支持从Excel表格批量导入病历\n每一行代表一个病人的病历数据')
        excel_info.setWordWrap(True)
        excel_layout.addWidget(excel_info)

        self.excel_import_label = QLabel('未导入Excel数据')
        excel_layout.addWidget(self.excel_import_label)

        import_excel_btn = QPushButton('📊 从Excel导入')
        import_excel_btn.clicked.connect(self.import_from_excel)
        excel_layout.addWidget(import_excel_btn)

        excel_layout.addStretch()
        excel_tab.setLayout(excel_layout)
        import_tabs.addTab(excel_tab, '📊 从Excel导入')

        emr_layout.addWidget(import_tabs)

        # 统计信息
        self.count_label = QLabel('📌 已添加病历: 0')
        self.count_label.setFont(QFont('Arial', 10, QFont.Bold))
        emr_layout.addWidget(self.count_label)

        emr_group.setLayout(emr_layout)
        layout.addWidget(emr_group)

        # 按钮
        button_box = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel)
        button_box.button(QDialogButtonBox.Ok).setText('✅ 创建任务')
        button_box.button(QDialogButtonBox.Cancel).setText('❌ 取消')
        button_box.accepted.connect(self.validate_and_accept)
        button_box.rejected.connect(self.reject)
        layout.addWidget(button_box)

        self.setLayout(layout)

    def apply_styles(self):
        self.setStyleSheet("""
            QGroupBox {
                font-weight: bold;
                border: 2px solid #9b59b6;
                border-radius: 5px;
                margin-top: 10px;
                padding-top: 10px;
            }
            QPushButton {
                background-color: #8e44ad;
                color: white;
                border: none;
                border-radius: 5px;
                padding: 8px;
                font-weight: bold;
                min-height: 30px;
            }
            QPushButton:hover {
                background-color: #7d3c98;
            }
        """)

    def select_template(self):
        file_path, _ = QFileDialog.getOpenFileName(
            self, '选择Excel模板', '', 'Excel Files (*.xlsx *.xls)'
        )
        if file_path:
            self.template_path = file_path
            self.template_label.setText(f'✅ {os.path.basename(file_path)}')

    def select_output(self):
        file_path, _ = QFileDialog.getSaveFileName(
            self, '选择输出文件', '', 'Excel Files (*.xlsx)'
        )
        if file_path:
            if not file_path.endswith('.xlsx'):
                file_path += '.xlsx'
            self.output_path = file_path
            self.output_label.setText(f'✅ {os.path.basename(file_path)}')

    def add_emr_files(self):
        """添加病历文件"""
        file_paths, _ = QFileDialog.getOpenFileNames(
            self, '选择电子病历文件', '',
            'All Files (*.*);;'
            'Text Files (*.txt *.text);;'
            'Markdown Files (*.md *.markdown);;'
            'Word Files (*.docx *.doc);;'
            'PDF Files (*.pdf);;'
            'Rich Text (*.rtf);;'
            'HTML Files (*.html *.htm);;'
            'OpenDocument (*.odt);;'
            'XML Files (*.xml);;'
            'JSON Files (*.json);;'
            'CSV Files (*.csv);;'
            'Log Files (*.log)'
        )

        if file_paths:
            for file_path in file_paths:
                # 读取文件内容
                try:
                    content = self.read_file(file_path)
                    self.emr_data.append({
                        'source': os.path.basename(file_path),
                        'content': content
                    })
                    self.file_list.addItem(f'✅ {os.path.basename(file_path)}')
                except Exception as e:
                    QMessageBox.warning(self, '警告', f'读取文件失败 {file_path}:\n{str(e)}')

            self.update_count()

    def add_emr_folder(self):
        """添加病历文件夹"""
        folder_path = QFileDialog.getExistingDirectory(self, '选择文件夹')
        if folder_path:
            # 支持的文件扩展名
            extensions = [
                '.txt', '.text',           # 纯文本
                '.md', '.markdown',        # Markdown
                '.docx', '.doc',           # Word文档
                '.pdf',                    # PDF
                '.rtf',                    # Rich Text Format
                '.html', '.htm',           # HTML
                '.odt',                    # OpenDocument Text
                '.xml',                    # XML
                '.json',                   # JSON
                '.csv',                    # CSV
                '.log'                     # 日志文件
            ]
            count = 0

            for root, dirs, files in os.walk(folder_path):
                for file in files:
                    _, ext = os.path.splitext(file)
                    if ext.lower() in extensions:
                        file_path = os.path.join(root, file)
                        try:
                            content = self.read_file(file_path)
                            self.emr_data.append({
                                'source': file,
                                'content': content
                            })
                            count += 1
                        except Exception as e:
                            logger.warning(f"跳过文件 {file}: {str(e)}")

            self.file_list.addItem(f'✅ 从文件夹添加了 {count} 个文件')
            self.update_count()

    def import_from_excel(self):
        """从Excel导入"""
        dialog = ExcelImportDialog(self)
        if dialog.exec_() == QDialog.Accepted:
            excel_data = dialog.get_emr_data()
            if excel_data:
                self.emr_data.extend(excel_data)
                self.excel_import_label.setText(f'✅ 已从Excel导入 {len(excel_data)} 条病历')
                self.update_count()
                # 保存对话框引用以便访问分批设置
                self.excel_dialog = dialog

    def clear_file_list(self):
        """清空文件列表"""
        self.file_list.clear()
        self.emr_data = []
        self.excel_import_label.setText('未导入Excel数据')
        self.update_count()

    def update_count(self):
        """更新计数"""
        self.count_label.setText(f'📌 已添加病历: {len(self.emr_data)}')

    def read_file(self, file_path: str) -> str:
        """读取文件内容 - 支持多种格式"""
        _, ext = os.path.splitext(file_path)
        ext = ext.lower()

        try:
            # 纯文本文件
            if ext in ['.txt', '.text', '.log']:
                with open(file_path, 'r', encoding='utf-8') as f:
                    return f.read()

            # Markdown文件
            elif ext in ['.md', '.markdown']:
                with open(file_path, 'r', encoding='utf-8') as f:
                    content = f.read()
                # 可选：移除Markdown标记，只保留纯文本
                # 如果需要保留Markdown格式，直接返回content
                return content

            # Word文档 (.docx)
            elif ext == '.docx':
                try:
                    import docx
                    doc = docx.Document(file_path)
                    return '\n'.join([para.text for para in doc.paragraphs])
                except ImportError:
                    raise Exception('需要安装 python-docx: pip install python-docx')

            # 旧版Word文档 (.doc)
            elif ext == '.doc':
                try:
                    import textract
                    text = textract.process(file_path).decode('utf-8')
                    return text
                except ImportError:
                    raise Exception('需要安装 textract: pip install textract\n'
                                  '或使用 LibreOffice 转换为 .docx 格式')
                except Exception as e:
                    # 如果textract失败，尝试提示用户
                    raise Exception(f'读取.doc文件失败。建议：\n'
                                  f'1. 安装 textract: pip install textract\n'
                                  f'2. 或将文件转换为 .docx 格式\n'
                                  f'错误详情: {str(e)}')

            # PDF文件
            elif ext == '.pdf':
                try:
                    import PyPDF2
                    with open(file_path, 'rb') as f:
                        reader = PyPDF2.PdfReader(f)
                        text = ''
                        for page in reader.pages:
                            text += page.extract_text()
                        return text
                except ImportError:
                    raise Exception('需要安装 PyPDF2: pip install PyPDF2')

            # RTF文件
            elif ext == '.rtf':
                try:
                    from striprtf.striprtf import rtf_to_text
                    with open(file_path, 'r', encoding='utf-8') as f:
                        rtf_content = f.read()
                    return rtf_to_text(rtf_content)
                except ImportError:
                    raise Exception('需要安装 striprtf: pip install striprtf')
                except Exception as e:
                    # 尝试备用方法
                    logger.warning(f"RTF解析失败，尝试作为纯文本读取: {str(e)}")
                    with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                        return f.read()

            # HTML文件
            elif ext in ['.html', '.htm']:
                try:
                    from bs4 import BeautifulSoup
                    with open(file_path, 'r', encoding='utf-8') as f:
                        html_content = f.read()
                    soup = BeautifulSoup(html_content, 'html.parser')
                    # 提取纯文本
                    return soup.get_text(separator='\n', strip=True)
                except ImportError:
                    raise Exception('需要安装 beautifulsoup4: pip install beautifulsoup4')

            # OpenDocument Text (.odt)
            elif ext == '.odt':
                try:
                    from odfpy import opendocument
                    from odfpy.text import P
                    doc = opendocument.load(file_path)
                    paragraphs = doc.getElementsByType(P)
                    text_content = []
                    for para in paragraphs:
                        text_content.append(str(para))
                    return '\n'.join(text_content)
                except ImportError:
                    raise Exception('需要安装 odfpy: pip install odfpy')
                except Exception as e:
                    # 尝试使用textract作为备用方案
                    try:
                        import textract
                        text = textract.process(file_path).decode('utf-8')
                        return text
                    except:
                        raise Exception(f'读取.odt文件失败。建议安装: pip install odfpy\n错误: {str(e)}')

            # XML文件
            elif ext == '.xml':
                with open(file_path, 'r', encoding='utf-8') as f:
                    content = f.read()
                # 可选：解析XML并提取文本
                try:
                    from bs4 import BeautifulSoup
                    soup = BeautifulSoup(content, 'xml')
                    return soup.get_text(separator='\n', strip=True)
                except:
                    # 如果解析失败，返回原始内容
                    return content

            # JSON文件
            elif ext == '.json':
                with open(file_path, 'r', encoding='utf-8') as f:
                    json_data = json.load(f)
                # 将JSON转换为可读格式
                return json.dumps(json_data, indent=2, ensure_ascii=False)

            # CSV文件
            elif ext == '.csv':
                import csv
                with open(file_path, 'r', encoding='utf-8') as f:
                    reader = csv.reader(f)
                    rows = []
                    for row in reader:
                        rows.append(' | '.join(row))
                    return '\n'.join(rows)

            # 其他文件类型，尝试作为UTF-8文本读取
            else:
                try:
                    with open(file_path, 'r', encoding='utf-8') as f:
                        return f.read()
                except UnicodeDecodeError:
                    # 如果UTF-8失败，尝试其他编码
                    try:
                        with open(file_path, 'r', encoding='gbk') as f:
                            return f.read()
                    except:
                        with open(file_path, 'r', encoding='latin-1') as f:
                            return f.read()

        except Exception as e:
            # 详细的错误信息
            error_msg = f'读取文件失败: {os.path.basename(file_path)}\n'
            error_msg += f'文件类型: {ext}\n'
            error_msg += f'错误: {str(e)}'
            raise Exception(error_msg)

    def validate_and_accept(self):
        """验证并接受"""
        if not self.template_path:
            QMessageBox.warning(self, '错误', '请选择Excel模板')
            return

        if not self.output_path:
            QMessageBox.warning(self, '错误', '请选择输出路径')
            return

        if not self.emr_data:
            QMessageBox.warning(self, '错误', '请至少添加一个病历')
            return

        self.accept()

    def get_task_data(self):
        """获取任务数据"""
        return {
            'name': self.name_input.text(),
            'template_path': self.template_path,
            'output_path': self.output_path,
            'emr_data': self.emr_data
        }


# ==================== 错误报告对话框 ====================

class ErrorReportDialog(QDialog):
    """错误报告对话框"""
    def __init__(self, task: Task, parent=None):
        super().__init__(parent)
        self.task = task
        self.setWindowTitle(f'错误报告 - {task.name}')
        self.setMinimumSize(800, 600)
        self.init_ui()

    def init_ui(self):
        layout = QVBoxLayout()

        # 标题
        title = QLabel('⚠️ 错误统计与详情')
        title.setFont(QFont('Arial', 14, QFont.Bold))
        layout.addWidget(title)

        # 统计信息
        stats_group = QGroupBox('统计摘要')
        stats_layout = QFormLayout()

        stats_layout.addRow('总病历数:', QLabel(str(self.task.total)))
        stats_layout.addRow('成功数:', QLabel(f'✅ {self.task.success_count}'))
        stats_layout.addRow('失败数:', QLabel(f'❌ {self.task.failed_count}'))

        success_rate = (self.task.success_count / self.task.total * 100) if self.task.total > 0 else 0
        stats_layout.addRow('成功率:', QLabel(f'{success_rate:.1f}%'))

        stats_group.setLayout(stats_layout)
        layout.addWidget(stats_group)

        # 错误列表
        error_group = QGroupBox('错误详情')
        error_layout = QVBoxLayout()

        self.error_table = QTableWidget()
        self.error_table.setColumnCount(4)
        self.error_table.setHorizontalHeaderLabels(['来源', '错误类型', '错误信息', '时间'])
        self.error_table.horizontalHeader().setSectionResizeMode(2, QHeaderView.Stretch)

        # 填充错误数据
        self.error_table.setRowCount(len(self.task.errors))
        for i, error in enumerate(self.task.errors):
            self.error_table.setItem(i, 0, QTableWidgetItem(error['source']))
            self.error_table.setItem(i, 1, QTableWidgetItem(error['error_type']))
            self.error_table.setItem(i, 2, QTableWidgetItem(error['message']))
            self.error_table.setItem(i, 3, QTableWidgetItem(error['timestamp']))

        error_layout.addWidget(self.error_table)
        error_group.setLayout(error_layout)
        layout.addWidget(error_group)

        # 按钮
        btn_layout = QHBoxLayout()

        export_btn = QPushButton('💾 导出错误报告')
        export_btn.clicked.connect(self.export_report)
        btn_layout.addWidget(export_btn)

        close_btn = QPushButton('关闭')
        close_btn.clicked.connect(self.accept)
        btn_layout.addWidget(close_btn)

        layout.addLayout(btn_layout)

        self.setLayout(layout)

    def export_report(self):
        """导出错误报告"""
        file_path, _ = QFileDialog.getSaveFileName(
            self, '导出错误报告', '', 'JSON Files (*.json);;Text Files (*.txt)'
        )

        if file_path:
            try:
                if file_path.endswith('.json'):
                    with open(file_path, 'w', encoding='utf-8') as f:
                        json.dump(self.task.errors, f, indent=2, ensure_ascii=False)
                else:
                    with open(file_path, 'w', encoding='utf-8') as f:
                        f.write(f"错误报告 - {self.task.name}\n")
                        f.write("=" * 80 + "\n\n")
                        f.write(f"总病历数: {self.task.total}\n")
                        f.write(f"成功数: {self.task.success_count}\n")
                        f.write(f"失败数: {self.task.failed_count}\n\n")
                        f.write("详细错误:\n")
                        f.write("-" * 80 + "\n")

                        for error in self.task.errors:
                            f.write(f"\n来源: {error['source']}\n")
                            f.write(f"类型: {error['error_type']}\n")
                            f.write(f"信息: {error['message']}\n")
                            f.write(f"时间: {error['timestamp']}\n")
                            f.write("-" * 80 + "\n")

                QMessageBox.information(self, '成功', f'错误报告已导出到:\n{file_path}')

            except Exception as e:
                QMessageBox.warning(self, '错误', f'导出失败:\n{str(e)}')


# ==================== 主窗口（续） ====================

class MainWindow(QMainWindow):
    """主窗口 - 商业化版本"""
    def __init__(self):
        super().__init__()
        self.task_queue = TaskQueue()
        self.task_queue.load_from_file()
        self.process_thread = None
        self.init_ui()
        self.apply_modern_style()
        self.add_entrance_animation()

    def init_ui(self):
        self.setWindowTitle('医疗电子病历数据提取系统 Pro v2.1')
        self.setGeometry(100, 100, 1400, 900)

        # 创建中心部件
        central_widget = QWidget()
        self.setCentralWidget(central_widget)

        # 主布局
        main_layout = QVBoxLayout()

        # 顶部信息栏
        header = QLabel('🏥 医疗电子病历数据提取系统 Pro')
        header.setFont(QFont('Arial', 18, QFont.Bold))
        header.setStyleSheet('color: #2c3e50; padding: 10px;')
        main_layout.addWidget(header)

        # 创建标签页
        tab_widget = QTabWidget()
        tab_widget.setTabPosition(QTabWidget.North)

        # API配置标签页
        self.api_config_widget = ModernAPIConfigWidget()
        tab_widget.addTab(self.api_config_widget, 'API配置')

        # 模板配置标签页
        self.template_config_widget = TemplateConfigWidget()
        tab_widget.addTab(self.template_config_widget, '模板配置')

        # 任务队列标签页（使用改进版）
        self.task_queue_widget = ModernTaskQueueWidget(self.task_queue, self)
        tab_widget.addTab(self.task_queue_widget, '任务队列')

        main_layout.addWidget(tab_widget)

        # 状态栏
        self.status_bar = self.statusBar()
        self.status_bar.showMessage('✅ 系统就绪')

        # 进度条
        self.progress_bar = QProgressBar()
        self.progress_bar.setVisible(False)
        self.status_bar.addPermanentWidget(self.progress_bar)

        central_widget.setLayout(main_layout)

    def apply_modern_style(self):
        """应用现代化高级样式"""
        self.setStyleSheet("""""")
        # self.setStyleSheet("""
        #     QMainWindow {
        #         background: qlineargradient(x1:0, y1:0, x2:1, y2:1,
        #             stop:0 #f5f7fa, stop:1 #e8eaf6);
        #     }
        #
        #     /* 标签页样式 */
        #     QTabWidget::pane {
        #         border: none;
        #         background-color: transparent;
        #         border-radius: 12px;
        #     }
        #
        #     QTabBar::tab {
        #         background: transparent;
        #         color: #64748b;
        #         padding: 14px 32px;
        #         margin-right: 8px;
        #         border: none;
        #         border-radius: 10px;
        #         font-size: 14px;
        #         font-weight: 600;
        #         transition: all 0.3s ease;
        #         min-width: 100px;
        #     }
        #
        #     QTabBar::tab:selected {
        #         background: qlineargradient(x1:0, y1:0, x2:1, y2:0,
        #             stop:0 #667eea, stop:1 #764ba2);
        #         color: white;
        #         transform: translateY(-2px);
        #     }
        #
        #     QTabBar::tab:hover:!selected {
        #         background: rgba(102, 126, 234, 0.1);
        #         color: #667eea;
        #     }
        #
        #     /* 分组框样式 */
        #     QGroupBox {
        #         background-color: white;
        #         border: none;
        #         border-radius: 12px;
        #         margin-top: 20px;
        #         padding-top: 28px;
        #         font-size: 15px;
        #         font-weight: 600;
        #         color: #1e293b;
        #     }
        #
        #     QGroupBox::title {
        #         subcontrol-origin: margin;
        #         left: 20px;
        #         padding: 0 10px;
        #         color: #667eea;
        #     }
        #
        #     /* 输入框样式 */
        #     QLineEdit, QTextEdit, QSpinBox, QDoubleSpinBox, QComboBox {
        #         background-color: #f8fafc;
        #         border: 2px solid #e2e8f0;
        #         border-radius: 8px;
        #         padding: 10px 14px;
        #         font-size: 13px;
        #         color: #1e293b;
        #         selection-background-color: #667eea;
        #     }
        #
        #     QLineEdit:focus, QTextEdit:focus, QSpinBox:focus,
        #     QDoubleSpinBox:focus, QComboBox:focus {
        #         border: 2px solid #667eea;
        #         background-color: white;
        #     }
        #
        #     QLineEdit:hover, QTextEdit:hover, QSpinBox:hover,
        #     QDoubleSpinBox:hover, QComboBox:hover {
        #         border: 2px solid #a5b4fc;
        #     }
        #
        #     /* 表格样式 */
        #     QTableWidget {
        #         background-color: white;
        #         border: none;
        #         border-radius: 12px;
        #         gridline-color: #f1f5f9;
        #         font-size: 10px;
        #     }
        #
        #     QTableWidget::item {
        #         padding: 12px 16px;
        #         border-bottom: 1px solid #f1f5f9;
        #     }
        #
        #     QTableWidget::item:selected {
        #         background: qlineargradient(x1:0, y1:0, x2:1, y2:0,
        #             stop:0 rgba(102, 126, 234, 0.1),
        #             stop:1 rgba(118, 75, 162, 0.1));
        #         color: #667eea;
        #     }
        #
        #     QHeaderView::section {
        #         background: qlineargradient(x1:0, y1:0, x2:0, y2:1,
        #             stop:0 #f8fafc, stop:1 #f1f5f9);
        #         color: #475569;
        #         padding: 12px;
        #         border: none;
        #         border-bottom: 2px solid #e2e8f0;
        #         font-weight: 600;
        #         font-size: 13px;
        #     }
        #
        #     /* 滚动条样式 */
        #     QScrollBar:vertical {
        #         background: #f8fafc;
        #         width: 10px;
        #         border-radius: 5px;
        #     }
        #
        #     QScrollBar::handle:vertical {
        #         background: #cbd5e1;
        #         border-radius: 5px;
        #         min-height: 30px;
        #     }
        #
        #     QScrollBar::handle:vertical:hover {
        #         background: #94a3b8;
        #     }
        #
        #     QScrollBar:horizontal {
        #         background: #f8fafc;
        #         height: 10px;
        #         border-radius: 5px;
        #     }
        #
        #     QScrollBar::handle:horizontal {
        #         background: #cbd5e1;
        #         border-radius: 5px;
        #         min-width: 30px;
        #     }
        #
        #     QScrollBar::handle:horizontal:hover {
        #         background: #94a3b8;
        #     }
        #
        #     /* 进度条样式 */
        #     QProgressBar {
        #         background-color: #f1f5f9;
        #         border: none;
        #         border-radius: 8px;
        #         height: 8px;
        #         text-align: center;
        #     }
        #
        #     QProgressBar::chunk {
        #         background: qlineargradient(x1:0, y1:0, x2:1, y2:0,
        #             stop:0 #667eea, stop:1 #764ba2);
        #         border-radius: 8px;
        #     }
        #
        #     /* 复选框样式 */
        #     QCheckBox {
        #         spacing: 8px;
        #         font-size: 13px;
        #         color: #475569;
        #     }
        #
        #     QCheckBox::indicator {
        #         width: 18px;
        #         height: 18px;
        #     }
        #
        #     /* 列表样式 */
        #     QListWidget {
        #         background-color: white;
        #         border: 1px solid #e2e8f0;
        #         border-radius: 8px;
        #         padding: 8px;
        #     }
        #
        #     QListWidget::item {
        #         padding: 10px;
        #         border-radius: 6px;
        #         margin: 2px;
        #     }
        #
        #     QListWidget::item:selected {
        #         background: qlineargradient(x1:0, y1:0, x2:1, y2:0,
        #             stop:0 #667eea, stop:1 #764ba2);
        #         color: white;
        #     }
        #
        #     QListWidget::item:hover:!selected {
        #         background: #f1f5f9;
        #     }
        #
        #     /* 状态栏样式 */
        #     QStatusBar {
        #         background: qlineargradient(x1:0, y1:0, x2:1, y2:0,
        #             stop:0 #1e293b, stop:1 #334155);
        #         color: white;
        #         font-weight: 500;
        #         font-size: 13px;
        #         padding: 8px;
        #     }
        #
        #     /* 标签样式 */
        #     QLabel {
        #         color: #334155;
        #         font-size: 13px;
        #     }
        # """)

    def add_entrance_animation(self):
        """添加窗口入场动画"""
        # 淡入效果
        self.setWindowOpacity(0)
        self.fade_animation = QPropertyAnimation(self, b"windowOpacity")
        self.fade_animation.setDuration(600)
        self.fade_animation.setStartValue(0)
        self.fade_animation.setEndValue(1)
        self.fade_animation.setEasingCurve(QEasingCurve.OutCubic)
        self.fade_animation.start()

        # 缩放效果
        original_geometry = self.geometry()
        start_geometry = QRect(
            original_geometry.x() + 50,
            original_geometry.y() + 50,
            original_geometry.width() - 100,
            original_geometry.height() - 100
        )
        self.setGeometry(start_geometry)

        self.scale_animation = QPropertyAnimation(self, b"geometry")
        self.scale_animation.setDuration(600)
        self.scale_animation.setStartValue(start_geometry)
        self.scale_animation.setEndValue(original_geometry)
        self.scale_animation.setEasingCurve(QEasingCurve.OutCubic)
        self.scale_animation.start()

    def start_task_processing(self, task):
        """开始处理任务"""
        # 验证API配置
        api_config = self.api_config_widget.get_config()
        if not api_config['api_key']:
            QMessageBox.warning(self, '错误', '❌ 请先在API配置中设置API Key')
            return

        # 获取模板配置
        template_config = self.template_config_widget.get_template_config()
        if not template_config['fields']:
            QMessageBox.warning(self, '错误', '❌ 模板中没有字段，请先配置模板')
            return

        # 创建处理线程
        self.process_thread = BatchProcessThread(task, api_config, template_config)
        self.process_thread.progress_updated.connect(self.on_progress_updated)
        self.process_thread.task_completed.connect(self.on_task_completed)
        self.process_thread.task_failed.connect(self.on_task_failed)
        self.process_thread.emr_processed.connect(self.on_emr_processed)
        self.process_thread.error_occurred.connect(self.on_error_occurred)

        # 显示进度条
        self.progress_bar.setVisible(True)
        self.progress_bar.setMaximum(task.total)
        self.progress_bar.setValue(0)

        # 更新UI
        self.task_queue_widget.start_btn.setEnabled(False)
        self.task_queue_widget.pause_btn.setEnabled(True)

        # 开始处理
        self.process_thread.start()
        self.status_bar.showMessage('🔄 任务处理中...')
        logger.info(f"开始处理任务: {task.name}")

    def on_progress_updated(self, current, total, status):
        """更新进度 - 带平滑动画"""
        # 平滑进度条动画
        self.progress_animation = QPropertyAnimation(self.progress_bar, b"value")
        self.progress_animation.setDuration(300)
        self.progress_animation.setStartValue(self.progress_bar.value())
        self.progress_animation.setEndValue(current)
        self.progress_animation.setEasingCurve(QEasingCurve.OutCubic)
        self.progress_animation.start()

        self.status_bar.showMessage(f'🔄 {status} ({current}/{total})')

    def on_emr_processed(self, index, data, success):
        """单个病历处理完成"""
        if success:
            logger.info(f"第{index + 1}个病历处理成功")
        else:
            logger.warning(f"第{index + 1}个病历处理失败")

    def on_error_occurred(self, source, error_type, message):
        """错误发生"""
        logger.error(f"[{source}] {error_type}: {message}")

    def on_task_completed(self, task_id, results):
        """任务完成"""
        task = self.task_queue.get_task(task_id)
        if task:
            try:
                # 导出到Excel
                self.export_task_to_excel(task)

                # 显示完成信息
                success_rate = (task.success_count / task.total * 100) if task.total > 0 else 0

                msg = f"""✅ 任务完成！

总病历数: {task.total}
成功: {task.success_count}
失败: {task.failed_count}
成功率: {success_rate:.1f}%

结果已导出到:
{task.output_path}
"""

                if task.errors:
                    msg += f"\n⚠️ 有{len(task.errors)}个错误，是否查看错误报告？"

                    reply = QMessageBox.question(
                        self, '任务完成', msg,
                        QMessageBox.Yes | QMessageBox.No
                    )

                    if reply == QMessageBox.Yes:
                        # 显示错误报告
                        error_dialog = ErrorReportDialog(task, self)
                        error_dialog.exec_()
                else:
                    QMessageBox.information(self, '任务完成', msg)

                self.status_bar.showMessage(f'✅ 任务完成: {task.name}')
                logger.info(f"任务完成: {task.name}, 成功率: {success_rate:.1f}%")

            except Exception as e:
                QMessageBox.warning(self, '警告', f'⚠️ 任务完成但导出失败:\n{str(e)}')
                logger.error(f"导出失败: {str(e)}")

        # 重置UI
        self.progress_bar.setVisible(False)
        self.task_queue_widget.refresh_task_list()
        self.task_queue_widget.start_btn.setEnabled(False)
        self.task_queue_widget.pause_btn.setEnabled(False)
        self.task_queue.save_to_file()

        # 检查是否启用自动处理
        if self.task_queue_widget.auto_process_enabled:
            next_task = self.task_queue.get_next_task()
            if next_task:
                logger.info(f"自动开始处理下一个任务: {next_task.name}")
                # 延迟1秒后自动开始
                QTimer.singleShot(1000, lambda: self.start_task_processing(next_task))

    def on_task_failed(self, task_id, error_msg):
        """任务失败"""
        task = self.task_queue.get_task(task_id)
        if task:
            self.status_bar.showMessage(f'❌ 任务失败: {task.name}')
            QMessageBox.critical(self, '错误', f'❌ 任务处理失败:\n{error_msg}')
            logger.error(f"任务失败: {task.name}, 错误: {error_msg}")

        # 重置UI
        self.progress_bar.setVisible(False)
        self.task_queue_widget.refresh_task_list()
        self.task_queue_widget.start_btn.setEnabled(True)
        self.task_queue_widget.pause_btn.setEnabled(False)
        self.task_queue.save_to_file()

    def export_task_to_excel(self, task: Task):
        """导出任务结果到Excel（复用 engine.export_rows_to_excel）"""
        if not task.results:
            raise Exception('没有可导出的数据')
        try:
            export_rows_to_excel(task.results, task.template_path, task.output_path)
            logger.info(f"数据已导出到: {task.output_path}")
        except Exception as e:
            raise Exception(f'导出Excel失败: {str(e)}')

    def pause_task_processing(self):
        """暂停任务"""
        if self.process_thread:
            self.process_thread.stop()
            self.task_queue_widget.start_btn.setEnabled(True)
            self.task_queue_widget.pause_btn.setEnabled(False)
            self.status_bar.showMessage('⏸️ 任务已暂停')
            logger.info("任务已暂停")

    def closeEvent(self, event):
        """窗口关闭事件"""
        # 保存任务队列
        self.task_queue.save_to_file()

        # 如果有正在处理的任务
        if self.process_thread and self.process_thread.isRunning():
            reply = QMessageBox.question(
                self, '确认退出',
                '⚠️ 有任务正在处理中，确定要退出吗？\n（下次启动时可以继续处理）',
                QMessageBox.Yes | QMessageBox.No
            )

            if reply == QMessageBox.Yes:
                self.process_thread.stop()
                self.process_thread.wait()
                event.accept()
            else:
                event.ignore()
        else:
            event.accept()


# ==================== 任务队列Widget（现代化版本） ====================

class ModernTaskQueueWidget(QWidget):
    """现代化任务队列界面 - 支持自动处理"""
    def __init__(self, task_queue, main_window):
        super().__init__()
        self.task_queue = task_queue
        self.main_window = main_window
        self.auto_process_enabled = False  # 自动处理开关
        self.init_ui()
        self.apply_styles()

    def init_ui(self):
        layout = QVBoxLayout()

        # 标题和自动处理开关
        title_layout = QHBoxLayout()
        title = QLabel('任务队列管理')
        title.setFont(QFont('SF Pro Display', 18, QFont.Bold))
        title.setStyleSheet('color: #1e293b;')
        title_layout.addWidget(title)
        title_layout.addStretch()

        # 自动处理开关
        self.auto_process_checkbox = QCheckBox('自动处理队列')
        self.auto_process_checkbox.setFont(QFont('SF Pro Display', 13, QFont.Bold))
        self.auto_process_checkbox.stateChanged.connect(self.on_auto_process_changed)
        title_layout.addWidget(self.auto_process_checkbox)

        refresh_btn = ModernButton('刷新', color='info')
        refresh_btn.setMinimumWidth(90)
        refresh_btn.clicked.connect(self.refresh_task_list)
        title_layout.addWidget(refresh_btn)

        layout.addLayout(title_layout)

        # 批量选择按钮组
        batch_select_layout = QHBoxLayout()
        batch_label = QLabel('批量选择')
        batch_label.setStyleSheet('font-weight: 600; color: #64748b; font-size: 14px;')
        batch_select_layout.addWidget(batch_label)

        select_all_btn = ModernButton('全选', color='info')
        select_all_btn.setMinimumWidth(100)
        select_all_btn.clicked.connect(self.select_all_tasks)
        batch_select_layout.addWidget(select_all_btn)

        deselect_all_btn = ModernButton('取消', color='dark')
        deselect_all_btn.setMinimumWidth(100)
        deselect_all_btn.clicked.connect(self.deselect_all_tasks)
        batch_select_layout.addWidget(deselect_all_btn)

        invert_select_btn = ModernButton('反选', color='warning')
        invert_select_btn.setMinimumWidth(100)
        invert_select_btn.clicked.connect(self.invert_selection_tasks)
        batch_select_layout.addWidget(invert_select_btn)

        batch_select_layout.addStretch()
        layout.addLayout(batch_select_layout)

        # 任务列表
        self.task_table = QTableWidget()
        self.task_table.setColumnCount(7)
        self.task_table.setHorizontalHeaderLabels(['任务名称', '状态', '进度', '病历数', '成功', '失败', '创建时间'])
        self.task_table.horizontalHeader().setSectionResizeMode(0, QHeaderView.Stretch)
        self.task_table.horizontalHeader().setMinimumHeight(45)
        self.task_table.horizontalHeader().setDefaultAlignment(Qt.AlignLeft | Qt.AlignVCenter)
        self.task_table.setSelectionBehavior(QTableWidget.SelectRows)
        self.task_table.setSelectionMode(QTableWidget.MultiSelection)  # 支持多选
        self.task_table.itemSelectionChanged.connect(self.on_task_selected)
        layout.addWidget(self.task_table)

        # 按钮组 - 使用ModernButton
        btn_layout = QHBoxLayout()
        btn_layout.setSpacing(12)

        new_task_btn = ModernButton('新建任务', color='primary')
        new_task_btn.setMinimumWidth(110)
        new_task_btn.clicked.connect(self.create_new_task)
        btn_layout.addWidget(new_task_btn)

        self.start_btn = ModernButton('开始处理', color='success')
        self.start_btn.setMinimumWidth(110)
        self.start_btn.clicked.connect(self.start_processing)
        self.start_btn.setEnabled(False)
        btn_layout.addWidget(self.start_btn)

        self.pause_btn = ModernButton('暂停', color='warning')
        self.pause_btn.setMinimumWidth(90)
        self.pause_btn.clicked.connect(self.pause_processing)
        self.pause_btn.setEnabled(False)
        btn_layout.addWidget(self.pause_btn)

        delete_btn = ModernButton('删除', color='danger')
        delete_btn.setMinimumWidth(90)
        delete_btn.clicked.connect(self.delete_selected_tasks)
        btn_layout.addWidget(delete_btn)

        view_errors_btn = ModernButton('查看错误', color='warning')
        view_errors_btn.setMinimumWidth(110)
        view_errors_btn.clicked.connect(self.view_errors)
        btn_layout.addWidget(view_errors_btn)

        export_btn = ModernButton('导出', color='info')
        export_btn.setMinimumWidth(90)
        export_btn.clicked.connect(self.export_selected_tasks)
        btn_layout.addWidget(export_btn)

        btn_layout.addStretch()
        layout.addLayout(btn_layout)

        # 任务详情
        detail_group = QGroupBox('任务详情')
        detail_layout = QVBoxLayout()

        self.detail_text = QTextBrowser()
        self.detail_text.setMaximumHeight(150)
        detail_layout.addWidget(self.detail_text)

        detail_group.setLayout(detail_layout)
        layout.addWidget(detail_group)

        self.setLayout(layout)
        self.refresh_task_list()

    def apply_styles(self):
        """应用任务队列特定样式"""
        # 使用全局现代化样式，不需要额外样式
        pass

    def on_auto_process_changed(self, state):
        """自动处理开关改变"""
        self.auto_process_enabled = (state == Qt.Checked)
        if self.auto_process_enabled:
            logger.info("已启用自动处理队列")
            # 检查是否有待处理任务
            next_task = self.task_queue.get_next_task()
            if next_task:
                reply = QMessageBox.question(
                    self, '自动处理',
                    f'发现待处理任务 "{next_task.name}"，是否立即开始？',
                    QMessageBox.Yes | QMessageBox.No
                )
                if reply == QMessageBox.Yes:
                    self.main_window.start_task_processing(next_task)
        else:
            logger.info("已禁用自动处理队列")

    def refresh_task_list(self):
        """刷新任务列表"""
        self.task_table.setRowCount(len(self.task_queue.tasks))

        # 设置行高
        for i in range(len(self.task_queue.tasks)):
            self.task_table.setRowHeight(i, 50)

        for i, task in enumerate(self.task_queue.tasks):
            # 任务名称
            self.task_table.setItem(i, 0, QTableWidgetItem(task.name))

            # 状态
            status_map = {
                'pending': '等待中',
                'processing': '处理中',
                'completed': '已完成',
                'failed': '失败',
                'paused': '已暂停'
            }
            status_text = status_map.get(task.status, task.status)
            status_item = QTableWidgetItem(status_text)

            if task.status == 'completed':
                status_item.setForeground(QColor(34, 197, 94))  # 绿色
            elif task.status == 'failed':
                status_item.setForeground(QColor(239, 68, 68))  # 红色
            elif task.status == 'processing':
                status_item.setForeground(QColor(59, 130, 246))  # 蓝色

            self.task_table.setItem(i, 1, status_item)

            # 进度
            progress_text = f'{task.progress}/{task.total}'
            self.task_table.setItem(i, 2, QTableWidgetItem(progress_text))

            # 病历数
            self.task_table.setItem(i, 3, QTableWidgetItem(str(task.total)))

            # 成功数
            self.task_table.setItem(i, 4, QTableWidgetItem(str(task.success_count)))

            # 失败数
            self.task_table.setItem(i, 5, QTableWidgetItem(str(task.failed_count)))

            # 创建时间
            time_str = task.created_time.strftime('%Y-%m-%d %H:%M')
            self.task_table.setItem(i, 6, QTableWidgetItem(time_str))

        self.task_table.resizeColumnsToContents()

    def create_new_task(self):
        """创建新任务"""
        dialog = ModernNewTaskDialog(self.main_window)
        if dialog.exec_() == QDialog.Accepted:
            task_data = dialog.get_task_data()

            # 检查是否从Excel导入并启用了分批
            if hasattr(dialog, 'excel_dialog') and dialog.excel_dialog:
                if dialog.excel_dialog.is_batch_enabled():
                    # 分批创建任务
                    batch_size = dialog.excel_dialog.get_batch_size()
                    emr_data = task_data['emr_data']
                    total_batches = (len(emr_data) + batch_size - 1) // batch_size

                    for batch_idx in range(total_batches):
                        start_idx = batch_idx * batch_size
                        end_idx = min(start_idx + batch_size, len(emr_data))
                        batch_data = emr_data[start_idx:end_idx]

                        # 创建批次任务
                        batch_name = f"{task_data['name']}_批次{batch_idx + 1}"
                        batch_output = task_data['output_path'].replace('.xlsx', f'_批次{batch_idx + 1}.xlsx')

                        task = Task(
                            batch_name,
                            task_data['template_path'],
                            batch_output
                        )

                        # 添加该批次的病历数据
                        for emr in batch_data:
                            task.add_emr_data(emr['source'], emr['content'])

                        self.task_queue.add_task(task)

                    self.refresh_task_list()
                    self.task_queue.save_to_file()

                    QMessageBox.information(
                        self, '成功',
                        f'✅ 已创建 {total_batches} 个批次任务\n'
                        f'共 {len(emr_data)} 个病历\n'
                        f'每批次 {batch_size} 个'
                    )
                    logger.info(f"创建分批任务: {total_batches}批次, 共{len(emr_data)}个病历")
                    return

            # 创建单个任务（未启用分批）
            task = Task(
                task_data['name'],
                task_data['template_path'],
                task_data['output_path']
            )

            # 添加病历数据
            for emr in task_data['emr_data']:
                task.add_emr_data(emr['source'], emr['content'])

            self.task_queue.add_task(task)
            self.refresh_task_list()
            self.task_queue.save_to_file()

            QMessageBox.information(self, '成功', f'✅ 任务已创建\n共 {len(task_data["emr_data"])} 个病历')
            logger.info(f"创建任务: {task.name}, 病历数: {task.total}")

    def on_task_selected(self):
        """任务选中事件"""
        selected_rows = self.task_table.selectedIndexes()
        if not selected_rows:
            self.detail_text.clear()
            self.start_btn.setEnabled(False)
            return

        row = selected_rows[0].row()
        if row >= len(self.task_queue.tasks):
            return

        task = self.task_queue.tasks[row]

        # 显示任务详情
        success_rate = (task.success_count / task.total * 100) if task.total > 0 else 0

        detail_html = f"""
        <h3>📋 {task.name}</h3>
        <table border="0" cellpadding="5">
            <tr><td><b>任务ID:</b></td><td>{task.id}</td></tr>
            <tr><td><b>模板文件:</b></td><td>{os.path.basename(task.template_path)}</td></tr>
            <tr><td><b>输出路径:</b></td><td>{task.output_path}</td></tr>
            <tr><td><b>病历总数:</b></td><td>{task.total}</td></tr>
            <tr><td><b>当前进度:</b></td><td>{task.progress}/{task.total}</td></tr>
            <tr><td><b>成功数:</b></td><td style="color: green;">{task.success_count}</td></tr>
            <tr><td><b>失败数:</b></td><td style="color: red;">{task.failed_count}</td></tr>
            <tr><td><b>成功率:</b></td><td>{success_rate:.1f}%</td></tr>
            <tr><td><b>状态:</b></td><td>{task.status}</td></tr>
            <tr><td><b>创建时间:</b></td><td>{task.created_time.strftime('%Y-%m-%d %H:%M:%S')}</td></tr>
        </table>
        """

        if task.error_msg:
            detail_html += f'<p><b style="color: red;">错误信息:</b> {task.error_msg}</p>'

        self.detail_text.setHtml(detail_html)

        # 更新按钮状态
        if task.status in ['pending', 'paused']:
            self.start_btn.setEnabled(True)
            self.pause_btn.setEnabled(False)
        elif task.status == 'processing':
            self.start_btn.setEnabled(False)
            self.pause_btn.setEnabled(True)
        else:
            self.start_btn.setEnabled(False)
            self.pause_btn.setEnabled(False)

    def start_processing(self):
        """开始处理"""
        selected_rows = self.task_table.selectedIndexes()
        if not selected_rows:
            return

        row = selected_rows[0].row()
        if row >= len(self.task_queue.tasks):
            return

        task = self.task_queue.tasks[row]
        self.main_window.start_task_processing(task)

    def pause_processing(self):
        """暂停处理"""
        self.main_window.pause_task_processing()

    def delete_task(self):
        """删除任务"""
        selected_rows = self.task_table.selectedIndexes()
        if not selected_rows:
            QMessageBox.warning(self, '提示', '请先选择要删除的任务')
            return

        row = selected_rows[0].row()
        if row >= len(self.task_queue.tasks):
            return

        task = self.task_queue.tasks[row]

        if task.status == 'processing':
            QMessageBox.warning(self, '提示', '无法删除正在处理的任务')
            return

        reply = QMessageBox.question(
            self, '确认删除',
            f'确定要删除任务 "{task.name}" 吗？',
            QMessageBox.Yes | QMessageBox.No
        )

        if reply == QMessageBox.Yes:
            self.task_queue.remove_task(task.id)
            self.refresh_task_list()
            self.task_queue.save_to_file()
            logger.info(f"删除任务: {task.name}")

    def view_errors(self):
        """查看错误"""
        selected_rows = self.task_table.selectedIndexes()
        if not selected_rows:
            QMessageBox.warning(self, '提示', '请先选择任务')
            return

        row = selected_rows[0].row()
        if row >= len(self.task_queue.tasks):
            return

        task = self.task_queue.tasks[row]

        if not task.errors:
            QMessageBox.information(self, '提示', '✅ 该任务没有错误')
            return

        # 显示错误报告
        error_dialog = ErrorReportDialog(task, self)
        error_dialog.exec_()

    def export_task_result(self):
        """导出任务结果"""
        selected_rows = self.task_table.selectedIndexes()
        if not selected_rows:
            QMessageBox.warning(self, '提示', '请先选择要导出的任务')
            return

        row = selected_rows[0].row()
        if row >= len(self.task_queue.tasks):
            return

        task = self.task_queue.tasks[row]

        if task.status != 'completed':
            QMessageBox.warning(self, '提示', '只能导出已完成的任务')
            return

        try:
            self.main_window.export_task_to_excel(task)
            QMessageBox.information(self, '成功', f'✅ 结果已导出到:\n{task.output_path}')
        except Exception as e:
            QMessageBox.critical(self, '错误', f'❌ 导出失败:\n{str(e)}')

    # ==================== 批量操作方法 ====================

    def select_all_tasks(self):
        """全选所有任务"""
        self.task_table.selectAll()
        logger.info("已全选所有任务")

    def deselect_all_tasks(self):
        """取消全选"""
        self.task_table.clearSelection()
        logger.info("已取消选择所有任务")

    def invert_selection_tasks(self):
        """反选任务"""
        total_rows = self.task_table.rowCount()
        for row in range(total_rows):
            if self.task_table.item(row, 0).isSelected():
                self.task_table.item(row, 0).setSelected(False)
            else:
                self.task_table.selectRow(row)
        logger.info("已反选任务")

    def delete_selected_tasks(self):
        """批量删除选中的任务"""
        selected_rows = self.task_table.selectionModel().selectedRows()
        if not selected_rows:
            QMessageBox.warning(self, '提示', '请先选择要删除的任务')
            return

        # 获取选中的任务
        selected_tasks = []
        processing_tasks = []
        for index in selected_rows:
            row = index.row()
            if row < len(self.task_queue.tasks):
                task = self.task_queue.tasks[row]
                if task.status == 'processing':
                    processing_tasks.append(task.name)
                else:
                    selected_tasks.append(task)

        # 检查是否有正在处理的任务
        if processing_tasks:
            QMessageBox.warning(
                self, '提示',
                f'⚠️ 以下任务正在处理中，无法删除：\n' + '\n'.join(processing_tasks)
            )
            if not selected_tasks:
                return

        # 确认删除
        reply = QMessageBox.question(
            self, '确认删除',
            f'确定要删除选中的 {len(selected_tasks)} 个任务吗？',
            QMessageBox.Yes | QMessageBox.No
        )

        if reply == QMessageBox.Yes:
            # 删除选中的任务
            for task in selected_tasks:
                self.task_queue.remove_task(task.id)
                logger.info(f"删除任务: {task.name}")

            self.refresh_task_list()
            self.task_queue.save_to_file()

            QMessageBox.information(self, '成功', f'✅ 已删除 {len(selected_tasks)} 个任务')

    def export_selected_tasks(self):
        """批量导出选中的任务"""
        selected_rows = self.task_table.selectionModel().selectedRows()
        if not selected_rows:
            QMessageBox.warning(self, '提示', '请先选择要导出的任务')
            return

        # 获取选中的已完成任务
        completed_tasks = []
        incomplete_tasks = []

        for index in selected_rows:
            row = index.row()
            if row < len(self.task_queue.tasks):
                task = self.task_queue.tasks[row]
                if task.status == 'completed':
                    completed_tasks.append(task)
                else:
                    incomplete_tasks.append(task.name)

        if not completed_tasks:
            QMessageBox.warning(self, '提示', '所选任务中没有已完成的任务')
            return

        # 提示有未完成的任务
        if incomplete_tasks:
            reply = QMessageBox.question(
                self, '提示',
                f'⚠️ 以下任务未完成，将跳过：\n' + '\n'.join(incomplete_tasks[:5]) +
                (f'\n...等{len(incomplete_tasks)}个任务' if len(incomplete_tasks) > 5 else '') +
                f'\n\n继续导出 {len(completed_tasks)} 个已完成任务？',
                QMessageBox.Yes | QMessageBox.No
            )
            if reply == QMessageBox.No:
                return

        # 批量导出
        success_count = 0
        failed_tasks = []

        for task in completed_tasks:
            try:
                self.main_window.export_task_to_excel(task)
                success_count += 1
                logger.info(f"导出任务: {task.name}")
            except Exception as e:
                failed_tasks.append(f"{task.name}: {str(e)}")
                logger.error(f"导出任务失败 {task.name}: {str(e)}")

        # 显示结果
        msg = f'✅ 成功导出 {success_count}/{len(completed_tasks)} 个任务'
        if failed_tasks:
            msg += f'\n\n❌ 失败 {len(failed_tasks)} 个：\n' + '\n'.join(failed_tasks[:3])
            if len(failed_tasks) > 3:
                msg += f'\n...等{len(failed_tasks)}个'

        if failed_tasks:
            QMessageBox.warning(self, '导出完成', msg)
        else:
            QMessageBox.information(self, '成功', msg)


# ==================== 主函数 ====================

def main():
    app = QApplication(sys.argv)
    app.setStyle('Fusion')

    # 设置应用图标（如果有）
    # app.setWindowIcon(QIcon('icon.png'))

    window = MainWindow()
    window.show()

    logger.info("医疗电子病历数据提取系统 Pro v2.1 已启动")

    sys.exit(app.exec_())


if __name__ == '__main__':
    main()
